#!/usr/bin/env python3
"""Test helper: plays the HUMAN reviewer (never used by agents). Fills the review workbook for some rows and
writes an HTML-style export for others, so both import paths are exercised."""
import json, sys
from pathlib import Path
import openpyxl

d = Path(sys.argv[1])
al = json.loads((d / "alignment.json").read_text())
if len(sys.argv) > 2 and sys.argv[2] == "approve-all":
    json.dump({"format": "acord-alignment-review", "version": 1, "region": al["meta"]["region"], "slice": al["meta"]["scope"]["slice"],
               "alignmentGeneratedAt": al["meta"]["generatedAt"], "reviewer": "Priya (UK architect)",
               "entities": {e["id"]: {"decision": "approve", "basis": e["basis"]} for e in al["entities"]},
               "attributes": {}, "endpoints": {}}, open(d / "review-decisions.json", "w"), indent=1)
    print("reviewed: all", len(al["entities"]), "entities approved via HTML export")
    sys.exit(0)
xl_dec = {  # entity -> (decision, canonical name, comment)
    "Address": ("approve", "", ""), "Money": ("approve", "", ""), "Claim": ("approve", "", ""),
    "CreateClaimRequest": ("approve", "", "Command payload of Claim"), "UpdateClaimRequest": ("approve", "", ""),
    "FnolSubmission": ("extend", "", "FNOL is the first version of a Claim"), "ClaimStatus": ("extend", "", "Keep Draft/Rejected"),
    "Claimant": ("approve", "", ""), "ClaimDocument": ("approve", "", ""), "PolicyHolder": ("approve", "", ""),
    "ApprovalRequest": ("custom", "ClaimApproval", "No reference equivalent"),
    "FnolReceipt": ("reject", "", "Technical acknowledgement, not canonical"),
}
wb = openpyxl.load_workbook(d / "acord-alignment.xlsx")
ws = wb["Entities"]; hdr = [c.value for c in ws[1]]
for row in ws.iter_rows(min_row=2):
    rec = dict(zip(hdr, row))
    ent = rec["Entity"].value
    if ent in xl_dec:
        dec, cn, com = xl_dec[ent]
        rec["Decision"].value, rec["Canonical name"].value, rec["Comment"].value, rec["Reviewer"].value = dec, cn or None, com or None, "Asha (Claims architect)"
ws = wb["Attributes"]; hdr = [c.value for c in ws[1]]
for row in ws.iter_rows(min_row=2):
    rec = dict(zip(hdr, row))
    if rec["ID"].value == "Claim.createdBy":
        rec["Decision"].value, rec["Comment"].value, rec["Reviewer"].value = "exclude", "Audit field, not business data", "Asha (Claims architect)"
    if rec["ID"].value == "Claim.createdAt":
        rec["Decision"].value, rec["Canonical attribute name"].value, rec["Reviewer"].value = "rename", "creationTimestamp", "Asha (Claims architect)"
wb.save(d / "reviewed-claims.xlsx")
ents = {e["id"]: {"decision": "approve", "basis": e["basis"]} for e in al["entities"] if e["domain"] == "Policy" and e["name"] not in xl_dec}
eps = {o["id"]: {"decision": "include", "canonicalPath": "/claims/{claimId}/approval", "basis": o["basis"]}
       for o in al["endpoints"] if o["operationId"] == "ApproveClaim"}
json.dump({"format": "acord-alignment-review", "version": 1, "region": al["meta"]["region"], "slice": al["meta"]["scope"]["slice"],
           "alignmentGeneratedAt": al["meta"]["generatedAt"], "reviewer": "Marco (Policy architect)", "exportedAt": "now",
           "entities": ents, "attributes": {}, "endpoints": eps}, open(d / "review-decisions.json", "w"), indent=1)
print("reviewed:", len(xl_dec), "via Excel,", len(ents), "entities +", len(eps), "endpoints via HTML export")
