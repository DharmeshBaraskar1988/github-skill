#!/usr/bin/env python3
"""
regional-view: aggregate many applications (any region) into one HTML page and one Excel workbook.

Input folder (any depth). For every application it accepts either
  a) an analysis folder  (contains analysis-summary.json + 2-/3-/4- artifacts)   -> used as is
  b) a bare OpenAPI file (.yaml/.yml/.json with an 'openapi' key)               -> analysed inline
     region/application come from info.x-region / info.x-application, else from the
     folder layout <input>/<REGION>/<application>/<file>, else from the file name.
     A decisions.yaml next to the file is honoured.

Usage:
  python build_regional_view.py --input api-catalog --domains api-catalog/domains.json --out api-catalog/regional-view
  python build_regional_view.py --config api-catalog.config.yaml
Outputs: regional-view.html, regional-view.xlsx, regional-view-data.json
"""
from __future__ import annotations

import argparse
import datetime as dt
import html
import json
import os
import re
import sys
from pathlib import Path

import yaml

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent.parent / "api-analysis" / "scripts"))
from analysis_core import run_analysis, business_name  # noqa: E402


def read_struct(p: Path):
    t = p.read_text(encoding="utf-8")
    return json.loads(t) if p.suffix == ".json" else yaml.safe_load(t)


def find_inputs(root: Path, skip_dirs: list) -> list:
    found, analysis_dirs = [], set()
    for s in sorted(root.rglob("analysis-summary.json")):
        if any(sd in s.parts for sd in skip_dirs):
            continue
        d = s.parent
        analysis_dirs.add(d)
        found.append({"kind": "analysis", "dir": d})
    for f in sorted(list(root.rglob("*.yaml")) + list(root.rglob("*.yml")) + list(root.rglob("*.json"))):
        if any(sd in f.parts for sd in skip_dirs) or f.parent in analysis_dirs:
            continue
        if "discovery" in f.parts:       # discovery output is superseded by its analysis folder
            region_app = f.parent.relative_to(root).parts[-2:] if len(f.parent.relative_to(root).parts) >= 2 else ()
            if any(d.parts[-2:] == region_app for d in analysis_dirs):
                continue
        if f.name in ("decisions.yaml", "overrides.yaml", "domains.json", "validation.json", "inventory.json",
                      "regional-view-data.json") or f.name.endswith(("-capabilities.json", "-attributes.json", "-report.json")):
            continue
        try:
            doc = read_struct(f)
        except Exception:  # noqa: BLE001
            continue
        if isinstance(doc, dict) and str(doc.get("openapi", "")).startswith("3") and doc.get("paths") is not None:
            found.append({"kind": "spec", "file": f, "doc": doc})
    return found


def infer_region_app(root: Path, f: Path, doc: dict):
    info = doc.get("info", {}) or {}
    region, app = info.get("x-region"), info.get("x-application")
    rel = f.relative_to(root).parts
    if not region and len(rel) >= 3:
        region = rel[-3]
    if not app and len(rel) >= 2:
        app = rel[-2]
    if not app:
        app = re.sub(r"[-_.]?(openapi|swagger|spec)$", "", f.stem, flags=re.I) or f.stem
    return (region or "UNSPECIFIED"), app


def load_app(root: Path, item: dict, domains: dict, cfg: dict) -> dict:
    if item["kind"] == "analysis":
        d = item["dir"]
        summ = read_struct(d / "analysis-summary.json")
        dc = read_struct(d / "2-domains-capabilities.json")
        ea = read_struct(d / "3-entities-attributes.json")
        du = read_struct(d / "4-duplicates-report.json")
        val = read_struct(d / "validation.json") if (d / "validation.json").exists() else None
        disc_val = None
        dv = d.parent.parent.parent / "discovery" / summ["region"] / summ["application"] / "validation.json"
        if dv.exists():
            disc_val = read_struct(dv)
        return {"region": summ["region"], "application": summ["application"], "mode": "analysis",
                "source": str(d.relative_to(root)), "summary": summ, "dc": dc, "ea": ea, "du": du,
                "validation": {"analysis": val["status"] if val else "missing",
                               "discovery": disc_val["status"] if disc_val else "missing",
                               "errors": (val or {}).get("errors", []), "warnings": (val or {}).get("warnings", [])}}
    f, doc = item["file"], item["doc"]
    region, app = infer_region_app(root, f, doc)
    doc.setdefault("info", {})
    doc["info"]["x-region"], doc["info"]["x-application"] = region, app
    dec_p = f.parent / "decisions.yaml"
    decisions = yaml.safe_load(dec_p.read_text(encoding="utf-8")) if dec_p.exists() else {}
    res = run_analysis(doc, domains, decisions or {}, cfg)
    return {"region": region, "application": app, "mode": "inline-analysis", "source": str(f.relative_to(root)),
            "summary": res["summary"], "dc": res["domainsCapabilities"], "ea": res["entitiesAttributes"],
            "du": res["duplicates"],
            "validation": {"analysis": "not-run (inline)", "discovery": "unknown (external spec)", "errors": [], "warnings": []}}


def cross_app_duplicates(entities: list, threshold: float) -> list:
    out = []
    objs = [e for e in entities if e["kind"] == "object" and len(e["attrNames"]) >= 2]
    for i, a in enumerate(objs):
        for b in objs[i + 1:]:
            if a["appId"] == b["appId"]:
                continue
            pa, pb = set(a["attrNames"]), set(b["attrNames"])
            inter = pa & pb
            same_name = business_name(a["name"]).lower() == business_name(b["name"]).lower()
            jac = len(inter) / len(pa | pb) if pa | pb else 0
            cont = len(inter) / min(len(pa), len(pb)) if pa and pb else 0
            if pa == pb:
                match = "exact"
            elif same_name and jac >= 0.4:
                match = "same-name"
            elif cont >= threshold and len(inter) >= 3:
                match = "near"
            else:
                continue
            out.append({"region": a["region"] if a["region"] == b["region"] else "cross-region",
                        "a": {"appId": a["appId"], "entity": a["name"]}, "b": {"appId": b["appId"], "entity": b["name"]},
                        "match": match, "jaccard": round(jac, 2), "containment": round(cont, 2),
                        "shared": sorted(inter), "onlyA": sorted(pa - pb), "onlyB": sorted(pb - pa)})
    return out


def build_data(root: Path, domains: dict, cfg: dict, skip_dirs: list) -> dict:
    items = find_inputs(root, skip_dirs)
    apps, seen = [], set()
    for it in items:
        a = load_app(root, it, domains, cfg)
        key = (a["region"], a["application"])
        if key in seen:
            a["application"] = f"{a['application']}~{len(seen)}"
            key = (a["region"], a["application"])
        seen.add(key)
        apps.append(a)

    data = {"generatedAt": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
            "inputFolder": str(root), "regions": sorted({a["region"] for a in apps}),
            "apps": [], "domains": {}, "capabilities": [], "operations": [], "entities": [],
            "relations": [], "duplicates": [], "crossAppDuplicates": []}
    cat = {d["name"]: d for d in domains.get("domains", [])}
    for a in apps:
        app_id = f"{a['region']}/{a['application']}"
        s = a["summary"]
        data["apps"].append({"id": app_id, "region": a["region"], "application": a["application"], "mode": a["mode"],
                             "source": a["source"], "validation": a["validation"],
                             "counts": {"operations": s["operations"], "domains": len(s["domainsUsed"]),
                                        "capabilities": s["capabilities"], "entities": s["entities"],
                                        "relations": s["relations"], "duplicateGroups": s["duplicateGroups"],
                                        "unclassified": s["unclassifiedOperations"]},
                             "schemaToEntity": a["ea"].get("schemaToEntity", {}),
                             "excludedSchemas": a["ea"].get("excludedSchemas", []),
                             "duplicateMapping": a["du"].get("mapping", {})})
        for dd in a["dc"]["domains"]:
            dom = data["domains"].setdefault(dd["domain"], {
                "name": dd["domain"], "description": dd.get("description") or cat.get(dd["domain"], {}).get("description", ""),
                "inCatalogue": dd["domain"] in cat,
                "catalogueCapabilities": [c["name"] for c in cat.get(dd["domain"], {}).get("capabilities", [])],
                "apps": []})
            if app_id not in dom["apps"]:
                dom["apps"].append(app_id)
            for c in dd["capabilities"]:
                data["capabilities"].append({"appId": app_id, "domain": dd["domain"], "name": c["name"],
                                             "status": c["status"], "description": c.get("description", ""),
                                             "operations": [o["operationId"] for o in c["operations"]]})
        for o in a["dc"]["operations"]:
            data["operations"].append({"appId": app_id, "operationId": o["operationId"], "method": o["method"],
                                       "path": o["path"], "summary": o.get("summary", ""),
                                       "domain": o.get("domain") or "Unclassified", "capability": o.get("capability") or "",
                                       "confidence": o.get("confidence"), "assignment": o.get("assignment"),
                                       "sourceProject": o.get("sourceProject", ""), "sourceFile": o.get("sourceFile", ""),
                                       "request": o.get("request"), "responses": o.get("responses", {}),
                                       "parameters": o.get("parameters", []), "security": o.get("security", [])})
        for e in a["ea"]["entities"]:
            data["entities"].append({"appId": app_id, "region": a["region"], "name": e["name"], "kind": e["kind"],
                                     "domain": e.get("domain") or "", "description": e.get("description", ""),
                                     "roles": e["roles"], "schemaNames": e["schemaNames"],
                                     "sourceProjects": e.get("sourceProjects", []), "enumValues": e.get("enumValues"),
                                     "attributes": [{k: x.get(k) for k in ("name", "type", "format", "refEntity", "isCollection",
                                                                          "required", "nullable", "description", "enum", "presentIn")}
                                                    for x in e["attributes"]],
                                     "attrNames": [x["name"].lower() for x in e["attributes"]],
                                     "tree": e.get("tree"),
                                     "usedBy": [f"{u['method']} {u['path']}" for u in e.get("usedBy", [])],
                                     "usage": [{"operationId": u["operationId"], "roles": u.get("roles", []),
                                                "direct": u.get("direct", True)} for u in e.get("usedBy", [])]})
        for r in a["ea"]["relations"]:
            data["relations"].append(dict(r, appId=app_id))
        for g in a["du"]["groups"]:
            data["duplicates"].append({"appId": app_id, "members": g["members"], "match": g["match"],
                                       "similarity": g["similarity"], "decision": g["decision"],
                                       "canonical": g.get("canonical"), "note": g.get("note", ""),
                                       "memberDetails": g.get("memberDetails", [])})
    thr = float((cfg.get("regionalView", {}) or {}).get("crossAppThreshold", 0.8))
    data["crossAppDuplicates"] = cross_app_duplicates(data["entities"], thr)
    data["domains"] = sorted(data["domains"].values(), key=lambda d: (d["name"] == "Unclassified", d["name"]))
    data["totals"] = {"regions": len(data["regions"]), "applications": len(data["apps"]),
                      "operations": len(data["operations"]), "domains": len([d for d in data["domains"] if d["name"] != "Unclassified"]),
                      "capabilities": len({(c["domain"], c["name"]) for c in data["capabilities"]}),
                      "entities": len(data["entities"]), "relations": len(data["relations"]),
                      "duplicateGroups": len(data["duplicates"]), "crossAppDuplicates": len(data["crossAppDuplicates"])}
    return data


# ----------------------------------------------------------------------------- Excel
def write_xlsx(data: dict, path: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    head_font = Font(color="FFFFFF", bold=True)

    def sheet(title, headers, rows, widths=None):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for r in rows:
            ws.append([("" if v is None else (", ".join(map(str, v)) if isinstance(v, (list, tuple)) else v)) for v in r])
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(headers, 1):
            w = (widths or {}).get(h)
            if not w:
                longest = max([len(str(h))] + [len(str(ws.cell(row=r, column=i).value or "")) for r in range(2, min(ws.max_row, 300) + 1)])
                w = min(max(10, longest + 2), 60)
            ws.column_dimensions[get_column_letter(i)].width = w
        return ws

    ws0 = wb.active
    ws0.title = "Summary"
    ws0.append(["API catalogue - regional view"])
    ws0["A1"].font = Font(size=14, bold=True)
    ws0.append(["Generated", data["generatedAt"]])
    ws0.append(["Regions", ", ".join(data["regions"])])
    ws0.append([])
    for k, v in data["totals"].items():
        ws0.append([k, v])
    ws0.column_dimensions["A"].width, ws0.column_dimensions["B"].width = 26, 60

    sheet("Applications", ["Region", "Application", "Mode", "Discovery validation", "Analysis validation", "Operations",
                           "Domains", "Capabilities", "Entities", "Duplicate groups", "Unclassified", "Source"],
          [[a["region"], a["application"], a["mode"], a["validation"]["discovery"], a["validation"]["analysis"],
            a["counts"]["operations"], a["counts"]["domains"], a["counts"]["capabilities"], a["counts"]["entities"],
            a["counts"]["duplicateGroups"], a["counts"]["unclassified"], a["source"]] for a in data["apps"]])
    # domain x app matrix
    app_ids = [a["id"] for a in data["apps"]]
    mat = []
    for d in data["domains"]:
        row = [d["name"]]
        for aid in app_ids:
            row.append(sum(1 for c in data["capabilities"] if c["domain"] == d["name"] and c["appId"] == aid) or "")
        mat.append(row)
    sheet("Domain_Matrix", ["Domain \\ Application"] + app_ids, mat)
    sheet("Domains_Capabilities", ["Region", "Application", "Domain", "Capability", "Status", "Operations", "Description"],
          [[c["appId"].split("/")[0], c["appId"].split("/", 1)[1], c["domain"], c["name"], c["status"],
            c["operations"], c["description"]] for c in data["capabilities"]])
    sheet("Endpoints", ["Region", "Application", "Method", "Path", "operationId", "Domain", "Capability", "Confidence",
                        "Summary", "Source project", "Source file"],
          [[o["appId"].split("/")[0], o["appId"].split("/", 1)[1], o["method"], o["path"], o["operationId"], o["domain"],
            o["capability"], o["confidence"], o["summary"], o["sourceProject"], o["sourceFile"]] for o in data["operations"]],
          {"Path": 48, "Summary": 50})
    sheet("Entities", ["Region", "Application", "Entity", "Kind", "Domain", "Roles", "Schemas", "Attributes", "Used by",
                       "Description"],
          [[e["appId"].split("/")[0], e["appId"].split("/", 1)[1], e["name"], e["kind"], e["domain"], e["roles"],
            e["schemaNames"], len(e["attributes"]) if e["kind"] == "object" else len(e["enumValues"] or []),
            len(e["usedBy"]), e["description"]] for e in data["entities"]], {"Description": 60})
    attr_rows = []
    for e in data["entities"]:
        if e["kind"] == "enum":
            for v in e["enumValues"] or []:
                attr_rows.append([e["appId"].split("/")[0], e["appId"].split("/", 1)[1], e["name"], v, "enum value", "", "", "", ""])
            continue
        for x in e["attributes"]:
            typ = (("list of " if x["isCollection"] else "") + x["refEntity"]) if x["refEntity"] else \
                (x["type"] + (f" ({x['format']})" if x.get("format") and "(" not in x["type"] else ""))
            attr_rows.append([e["appId"].split("/")[0], e["appId"].split("/", 1)[1], e["name"], x["name"], typ,
                              "yes" if x["required"] else "", "yes" if x["nullable"] else "", x.get("presentIn") or "",
                              x["description"]])
    sheet("Attributes", ["Region", "Application", "Entity", "Attribute", "Type", "Required", "Nullable", "Present in schemas",
                         "Description"], attr_rows, {"Description": 60})
    sheet("Relations", ["Region", "Application", "From", "Relation", "To", "Via", "Cardinality"],
          [[r["appId"].split("/")[0], r["appId"].split("/", 1)[1], r["from"], r["type"], r["to"], r["via"], r["cardinality"]]
           for r in data["relations"]])
    sheet("Duplicates", ["Region", "Application", "Members", "Match", "Similarity", "Decision", "Canonical", "Note", "Used by"],
          [[g["appId"].split("/")[0], g["appId"].split("/", 1)[1], g["members"], g["match"], g["similarity"], g["decision"],
            g["canonical"], g["note"], "; ".join(f"{m['schema']}: {', '.join(m['usedBy'])}" for m in g["memberDetails"])]
           for g in data["duplicates"]], {"Used by": 80})
    sheet("Cross_App_Duplicates", ["Scope", "App A", "Entity A", "App B", "Entity B", "Match", "Jaccard", "Containment",
                                   "Shared attributes", "Only in A", "Only in B"],
          [[x["region"], x["a"]["appId"], x["a"]["entity"], x["b"]["appId"], x["b"]["entity"], x["match"], x["jaccard"],
            x["containment"], x["shared"], x["onlyA"], x["onlyB"]] for x in data["crossAppDuplicates"]])
    vrows = []
    for a in data["apps"]:
        for e in a["validation"]["errors"]:
            vrows.append([a["id"], "error", e["code"], e["message"]])
        for w in a["validation"]["warnings"]:
            vrows.append([a["id"], "warning", w["code"], w["message"]])
    sheet("Validation", ["Application", "Severity", "Code", "Message"], vrows, {"Message": 100})
    wb.save(path)


# ----------------------------------------------------------------------------- HTML
def write_html(data: dict, path: Path, title: str):
    tpl = (HERE.parent / "assets" / "regional-view.template.html").read_text(encoding="utf-8")
    payload = json.dumps(data, separators=(",", ":")).replace("</", "<\\/")
    out = tpl.replace("{{TITLE}}", html.escape(title)).replace("/*__DATA__*/null", payload)
    path.write_text(out, encoding="utf-8")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--config")
    ap.add_argument("--input")
    ap.add_argument("--domains")
    ap.add_argument("--out")
    ap.add_argument("--title")
    a = ap.parse_args()
    cfg = (yaml.safe_load(open(a.config, encoding="utf-8")) or {}) if a.config else {}
    base = Path(a.config).resolve().parent if a.config else Path.cwd()
    rv = cfg.get("regionalView", {}) or {}

    def rp(p):
        p = Path(p)
        return p if p.is_absolute() else (base / p).resolve()

    inp = rp(a.input or rv.get("inputDir", "api-catalog"))
    out = rp(a.out or rv.get("outputDir", "api-catalog/regional-view"))
    dom = rp(a.domains or cfg.get("analysis", {}).get("domainsJson", "api-catalog/domains.json"))
    out.mkdir(parents=True, exist_ok=True)
    domains = json.loads(dom.read_text(encoding="utf-8"))
    skip = [out.name, "regional-view", "canonical", "alignment", "reference", "regions", "style-examples"]
    data = build_data(inp, domains, cfg, skip)
    data["inputFolder"] = os.path.relpath(inp, out)
    if not data["apps"]:
        sys.exit(f"No analysis folders or OpenAPI files found under {inp}")
    from sensitive_scan import find_policy_file, load_policy, redact_obj, summarise
    data, red = redact_obj(data, load_policy(find_policy_file(out)))
    data["redactions"] = {"count": len(red), "summary": summarise(red)}
    (out / "regional-view-data.json").write_text(json.dumps(data, indent=1), encoding="utf-8")
    write_html(data, out / "regional-view.html", a.title or rv.get("title") or "API Catalogue - Regional View")
    write_xlsx(data, out / "regional-view.xlsx")
    sys.path.insert(0, str(HERE.parent.parent / "canonical-model" / "scripts"))
    from spec_viewer import scan, describe, load as load_doc, build as build_viewer  # noqa: E402
    vdocs = [describe(p, load_doc(p), inp) for p in scan(inp)]
    vdocs = [d for d in vdocs if d["kind"] != "other"]
    build_viewer(vdocs, out / "spec-viewer.html", (a.title or rv.get("title") or "API Catalogue") + " - OpenAPI specs")
    print(json.dumps({"output": str(out), "apps": [x["id"] for x in data["apps"]], "totals": data["totals"]}, indent=2))


if __name__ == "__main__":
    main()
