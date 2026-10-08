#!/usr/bin/env python3
"""
ACORD (reference) alignment for a region or a slice of it.

Inputs
  regional-view-data.json   from the regional-view skill (entities already de-duplicated per app)
  reference.json            ACORD export normalised by load_reference.py
  baseline (optional)       canonical-model.json of an approved region (e.g. EU when aligning UK)
  alignment-overrides.yaml  agent decisions on matching (in the output folder)
  approvals.yaml            human review decisions (in the output folder, written by import_review.py)
  synonyms.yaml (optional)  organisation vocabulary

Outputs (output folder)
  alignment.json            everything below, machine readable
  acord-alignment.html      interactive review page (filters, side-by-side attribute matching, decisions, export)
  acord-alignment.xlsx      review workbook (decision columns with drop-downs)
  alignment-report.md       summary

Usage
  python align.py --config api-catalog/regions/EU.yaml                      # whole region
  python align.py --config api-catalog/regions/EU.yaml --apps claims        # only the claims application
  python align.py --config api-catalog/regions/EU.yaml --domains Claims --entities Claim
"""
from __future__ import annotations

import argparse
import datetime as dt
import html
import os
import json
import sys
from pathlib import Path

import yaml

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from align_core import run_alignment  # noqa: E402
from load_reference import load_json, flatten_inheritance  # noqa: E402

DECISIONS_ENTITY = ["approve", "use-as-is", "extend", "custom", "reject", "pending"]
DECISIONS_ATTR = ["accept", "rename", "exclude"]
DECISIONS_OP = ["include", "exclude"]


def load_cfg(path):
    if not path:
        return {}, Path.cwd()
    return (yaml.safe_load(open(path, encoding="utf-8")) or {}), Path(path).resolve().parent


def rp(base: Path, p):
    if p is None:
        return None
    p = Path(p)
    return p if p.is_absolute() else (base / p).resolve()


def resolve_inputs(a, cfg, base):
    acfg = cfg.get("alignment", {}) or {}
    region = a.region or cfg.get("region")
    scope_cfg = acfg.get("scope", {}) or {}
    apps = a.apps.split(",") if a.apps else (scope_cfg.get("apps") or [])
    domains = a.domains.split(",") if a.domains else (scope_cfg.get("domains") or [])
    entities = a.entities.split(",") if a.entities else (scope_cfg.get("entities") or [])
    data_p = rp(base, a.data or cfg.get("regionalViewData") or "api-catalog/regional-view/regional-view-data.json")
    ref_p = rp(base, a.reference or (cfg.get("reference") or {}).get("json"))
    base_p = rp(base, a.baseline if a.baseline is not None else cfg.get("baseline"))
    syn_p = rp(base, a.synonyms or cfg.get("synonyms"))
    slice_name = a.name or ("-".join(apps) if apps else "all") + (("-" + "-".join(domains)) if domains else "") + \
        (("-" + "-".join(entities)) if entities else "")
    out = rp(base, a.out or acfg.get("outputDir") or f"api-catalog/alignment/{region}")
    if (apps or domains or entities) and not a.out:
        out = out / slice_name
    return dict(region=region, apps=apps, domains=domains, entities=entities, data=data_p, reference=ref_p,
                baseline=base_p, synonyms=syn_p, out=out, slice=slice_name, cfg=acfg)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--config")
    ap.add_argument("--region")
    ap.add_argument("--data")
    ap.add_argument("--reference")
    ap.add_argument("--baseline")
    ap.add_argument("--synonyms")
    ap.add_argument("--apps")
    ap.add_argument("--domains")
    ap.add_argument("--entities")
    ap.add_argument("--name", help="name of the slice folder")
    ap.add_argument("--out")
    a = ap.parse_args()
    cfg, base = load_cfg(a.config)
    r = resolve_inputs(a, cfg, base)
    if not r["region"]:
        sys.exit("region is required (--region or config)")
    data = json.loads(r["data"].read_text(encoding="utf-8"))
    app_ids = [x["id"] for x in data["apps"] if x["region"] == r["region"]
               and (not r["apps"] or x["application"] in r["apps"] or x["id"] in r["apps"])]
    refs = []
    if r["baseline"]:
        refs.append(load_json(r["baseline"]))
        refs[-1]["entities"] = flatten_inheritance(refs[-1]["entities"])
    if r["reference"]:
        refs.append(json.loads(r["reference"].read_text(encoding="utf-8")))
    if not refs:
        sys.exit("No reference model: set reference.json (load_reference.py output) and/or a baseline")
    out = r["out"]
    out.mkdir(parents=True, exist_ok=True)
    ov_p, ap_p = out / "alignment-overrides.yaml", out / "approvals.yaml"
    overrides = yaml.safe_load(ov_p.read_text(encoding="utf-8")) if ov_p.exists() else {}
    approvals = yaml.safe_load(ap_p.read_text(encoding="utf-8")) if ap_p.exists() else {}
    vocab = yaml.safe_load(r["synonyms"].read_text(encoding="utf-8")) if r["synonyms"] and r["synonyms"].exists() else {}
    scope = {"region": r["region"], "appIds": app_ids, "apps": r["apps"], "domains": r["domains"], "entities": r["entities"],
             "slice": r["slice"]}
    res = run_alignment(data, refs, overrides or {}, approvals or {}, scope, r["cfg"], vocab or {})
    res["meta"] = {"generatedAt": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"), "region": r["region"],
                   "scope": scope, "regionalViewData": os.path.relpath(r["data"], out),
                   "baseline": os.path.relpath(r["baseline"], out) if r["baseline"] else None,
                   "thresholds": {k: r["cfg"].get(k) for k in ("fullMatchThreshold", "partialMatchThreshold", "ambiguityMargin",
                                                                 "baselinePreference") if r["cfg"].get(k) is not None},
                   "overridesFile": ov_p.name if ov_p.exists() else None, "approvalsFile": ap_p.name if ap_p.exists() else None}
    res["overrides"] = overrides or {}
    (out / "alignment.json").write_text(json.dumps(res, indent=1), encoding="utf-8")
    write_md(res, out / "alignment-report.md")
    write_xlsx(res, out / "acord-alignment.xlsx")
    write_html(res, out / "acord-alignment.html")
    s = res["summary"]
    print(json.dumps({"output": str(out), "scope": {k: scope[k] for k in ("region", "appIds", "domains", "entities")},
                      "references": [f"{x['source']}: {x['name']} {x['version']}" for x in res["references"]],
                      **{k: s[k] for k in ("entities", "full", "partial", "none", "ambiguous", "attributes",
                                            "attributesMatched", "typeConflicts", "endpoints", "alignmentPercent", "review")}},
                     indent=2))


# ----------------------------------------------------------------------------- Markdown
def write_md(res, path):
    s, m = res["summary"], res["meta"]
    L = [f"# Reference alignment - {m['region']} ({m['scope']['slice']})", "",
         "References: " + "; ".join(f"{x['source']} = {x['name']} {x['version']}" for x in res["references"]), "",
         f"Overall alignment **{s['alignmentPercent']}%** - entities {s['entities']} (full {s['full']}, partial {s['partial']}, "
         f"none {s['none']}), attributes matched {s['attributesMatched']}/{s['attributes']}, ambiguous {s['ambiguous']}", "",
         "## Domains", "", "| Domain | Entities | Full | Partial | None | Alignment | Endpoint alignment |", "|---|---|---|---|---|---|---|"]
    for d in res["domains"]:
        L.append(f"| {d['name']} | {d['entities']} | {d['full']} | {d['partial']} | {d['none']} | {d['alignmentPercent']}% | "
                 f"{'-' if d.get('endpointAlignmentPercent') is None else str(d['endpointAlignmentPercent']) + '%'} |")
    L += ["", "## Entities", "", "| Entity | Apps | Best match | Alignment | Status | Recommendation | Proposed canonical |",
          "|---|---|---|---|---|---|---|"]
    for e in sorted(res["entities"], key=lambda x: (x["domain"], x["name"])):
        bm = e["match"]["reference"] + f" ({e['referenceSource']})" if e["match"] else "-"
        L.append(f"| {e['name']}{' (ambiguous)' if e['ambiguous'] else ''} | {', '.join(e['apps'])} | {bm} | {e['alignmentPercent']}% | "
                 f"{e['status']} | {e['recommendationLabel']} | {e['proposedCanonicalName']} |")
    L += ["", "## Endpoints", "", "| Method | Path | Entities | Alignment | Proposed canonical path |", "|---|---|---|---|---|"]
    for o in res["endpoints"]:
        L.append(f"| {o['method']} | `{o['path']}` | {', '.join(o['entities'])} | "
                 f"{'-' if o['alignmentPercent'] is None else str(o['alignmentPercent']) + '%'} | `{o['proposedPath']}` |")
    path.write_text("\n".join(L) + "\n", encoding="utf-8")


# ----------------------------------------------------------------------------- Excel
def write_xlsx(res, path):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule, DataBarRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    HEAD = PatternFill("solid", fgColor="1F3A5F")
    DEC = PatternFill("solid", fgColor="FFF4CC")
    DEC_HEAD = PatternFill("solid", fgColor="B07D00")
    st = res["review"]["state"]
    wb = Workbook()

    def sheet(title, headers, rows, decision_cols=(), dropdowns=None, hidden=(), widths=None, pct_cols=()):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for r in rows:
            ws.append([("" if v is None else (", ".join(map(str, v)) if isinstance(v, (list, tuple)) else v)) for v in r])
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=i)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = DEC_HEAD if h in decision_cols else HEAD
            c.alignment = Alignment(wrap_text=True, vertical="center")
            col = get_column_letter(i)
            if h in decision_cols:
                for rr in range(2, ws.max_row + 1):
                    ws.cell(row=rr, column=i).fill = DEC
            w = (widths or {}).get(h)
            if not w:
                longest = max([len(str(h))] + [len(str(ws.cell(row=rr, column=i).value or "")) for rr in range(2, min(ws.max_row, 200) + 1)])
                w = min(max(10, longest + 2), 55)
            ws.column_dimensions[col].width = w
            if h in hidden:
                ws.column_dimensions[col].hidden = True
            if h in pct_cols and ws.max_row > 1:
                rng = f"{col}2:{col}{ws.max_row}"
                ws.conditional_formatting.add(rng, DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color="4F8FD8"))
            if dropdowns and h in dropdowns and ws.max_row > 1:
                dv = DataValidation(type="list", formula1='"' + ",".join(dropdowns[h]) + '"', allow_blank=True)
                dv.add(f"{col}2:{col}{max(ws.max_row, 2) + 200}")
                ws.add_data_validation(dv)
        ws.freeze_panes = "C2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        return ws

    ws = wb.active
    ws.title = "README"
    m, s = res["meta"], res["summary"]
    lines = [
        (f"Reference alignment - {m['region']} ({m['scope']['slice']})", True),
        (f"Generated {m['generatedAt']}. References: " + "; ".join(f"{x['source']}: {x['name']} {x['version']}" for x in res["references"]), False),
        (f"Overall alignment {s['alignmentPercent']}% - {s['entities']} entities: {s['full']} full, {s['partial']} partial, {s['none']} none.", False),
        ("", False),
        ("How to review", True),
        ("1. Entities sheet: fill the yellow columns. Decision:", False),
        ("     approve   = accept the Recommendation column as it is", False),
        ("     use-as-is = adopt the reference entity unchanged (reference names and types)", False),
        ("     extend    = reference entity + our extra attributes as extensions", False),
        ("     custom    = new canonical entity (no reference equivalent)", False),
        ("     reject    = not part of the canonical model (say why in Comment)", False),
        ("     pending   = not decided yet", False),
        ("   'Reference (override)' picks another candidate (see Candidates sheet). 'Canonical name' renames the canonical entity.", False),
        ("2. Attributes sheet (optional): accept (default) / rename (+ Canonical attribute name) / exclude.", False),
        ("3. Endpoints sheet (optional): include (default) / exclude, and an optional canonical path / operationId.", False),
        ("4. Put your name in Reviewer. Save the file and hand it back (or run import_review.py --xlsx <this file>).", False),
        ("Rows already decided show the current decision. Do not change the hidden ID / Basis columns.", False),
        ("APPROVE ALL: on the Entities sheet filter the rows you want (e.g. Status = full), type 'approve' in the first visible", False),
        ("Decision cell, select it down to the last row and press Ctrl+D; do the same for Reviewer. (The HTML page has an 'Approve all' button.)", False),
        ("Entities listing several sources in 'Merges with' share a reference entity and become ONE canonical entity.", False),
    ]
    for t, b in lines:
        ws.append([t])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=b, size=13 if b else 11)
    ws.column_dimensions["A"].width = 130

    sheet("Domains", ["Domain", "Entities", "Full", "Partial", "None", "Alignment %", "Endpoints", "Endpoint alignment %",
                      "Reference subject areas"],
          [[d["name"], d["entities"], d["full"], d["partial"], d["none"], d["alignmentPercent"], d.get("endpoints"),
            d.get("endpointAlignmentPercent"), d["referenceSubjectAreas"]] for d in res["domains"]],
          pct_cols=("Alignment %", "Endpoint alignment %"))
    sheet("Applications", ["Application", "Entities", "Full", "Partial", "None", "Alignment %"],
          [[d["name"], d["entities"], d["full"], d["partial"], d["none"], d["alignmentPercent"]] for d in res["apps"]],
          pct_cols=("Alignment %",))

    ent_rows = []
    for e in sorted(res["entities"], key=lambda x: (x["domain"], x["name"])):
        a = st["entities"].get(e["id"], {})
        ent_rows.append([e["id"], e["basis"], e["domain"], e["name"], e["kind"], e["apps"], e["roles"],
                         [f"{mm['appId']}:{','.join(mm['schemas'])}" for mm in e["members"]],
                         e["match"]["reference"] if e["match"] else "", e["referenceSource"] or "", e["alignmentPercent"],
                         e["status"], e["recommendationLabel"], e["proposedCanonicalName"], e["mergesWith"],
                         "yes" if e["ambiguous"] else "",
                         a.get("decision", ""), a.get("reference", ""), a.get("canonicalName", ""), a.get("reviewer", ""),
                         a.get("comment", ""), a.get("state", "pending")])
    sheet("Entities", ["ID", "Basis", "Domain", "Entity", "Kind", "Applications", "Roles", "Source schemas", "Best match",
                       "Source", "Alignment %", "Status", "Recommendation", "Proposed canonical name", "Merges with", "Ambiguous",
                       "Decision", "Reference (override)", "Canonical name", "Reviewer", "Comment", "Review state"],
          ent_rows, decision_cols=("Decision", "Reference (override)", "Canonical name", "Reviewer", "Comment"),
          dropdowns={"Decision": DECISIONS_ENTITY}, hidden=("ID", "Basis"), pct_cols=("Alignment %",),
          widths={"Source schemas": 40, "Comment": 40, "Recommendation": 26})

    att_rows = []
    for e in sorted(res["entities"], key=lambda x: (x["domain"], x["name"])):
        for at in e["attributes"]:
            mt = at["match"]
            x = st["attributes"].get(at["id"], {})
            typ = (("list of " if at["isCollection"] else "") + at["refEntity"]) if at.get("refEntity") else at["type"]
            att_rows.append([at["id"], at["basis"], e["name"], at["name"], typ, "yes" if at["required"] else "",
                             at["description"], [f"{s_['appId']}:{'/'.join(s_['schemas'])}.{s_['attribute']}" for s_ in at["sources"]],
                             mt.get("reference") or "", mt.get("referenceType", ""), mt.get("nameSimilarity"),
                             mt.get("typeCompatibility"), mt.get("status"), at["proposedCanonicalName"], at["recommendation"],
                             at.get("typeConflict", ""),
                             x.get("decision", ""), x.get("canonicalName", ""), x.get("reviewer", ""), x.get("comment", "")])
    sheet("Attributes", ["ID", "Basis", "Entity", "Attribute", "Type", "Required", "Description", "Sources",
                         "Reference attribute", "Reference type", "Name similarity", "Type compatibility", "Match", "Proposed canonical",
                         "Recommendation", "Type conflict between apps", "Decision", "Canonical attribute name", "Reviewer", "Comment"],
          att_rows, decision_cols=("Decision", "Canonical attribute name", "Reviewer", "Comment"),
          dropdowns={"Decision": DECISIONS_ATTR}, hidden=("ID", "Basis"), widths={"Sources": 45, "Description": 40})

    op_rows = []
    for o in res["endpoints"]:
        x = st["endpoints"].get(o["id"], {})
        op_rows.append([o["id"], o["basis"], o["appId"], o["domain"], o["capability"], o["method"], o["path"], o["entities"],
                        o["alignmentPercent"], o["proposedPath"], x.get("decision", ""), x.get("canonicalPath", ""),
                        x.get("canonicalOperationId", ""), x.get("reviewer", ""), x.get("comment", "")])
    sheet("Endpoints", ["ID", "Basis", "Application", "Domain", "Capability", "Method", "Path", "Entities", "Alignment %",
                        "Proposed canonical path", "Decision", "Canonical path", "Canonical operationId", "Reviewer", "Comment"],
          op_rows, decision_cols=("Decision", "Canonical path", "Canonical operationId", "Reviewer", "Comment"),
          dropdowns={"Decision": DECISIONS_OP}, hidden=("ID", "Basis"), pct_cols=("Alignment %",))

    cand_rows = []
    for e in res["entities"]:
        for i, c in enumerate(e["candidates"], 1):
            cand_rows.append([e["name"], i, c["reference"], c["source"], round(100 * c["score"]), c["nameSimilarity"], c["coverage"],
                              c.get("referenceCoverage"), c.get("typeConflicts")])
    sheet("Candidates", ["Entity", "Rank", "Reference entity", "Source", "Score %", "Name similarity", "Coverage of our attributes",
                         "Coverage of reference attributes", "Type conflicts"], cand_rows, pct_cols=("Score %",))

    for wsn in ("Entities", "Endpoints"):
        w = wb[wsn]
        col = [c.value for c in w[1]].index("Alignment %" if wsn != "Entities" else "Alignment %") + 1
        L = get_column_letter(col)
        if w.max_row > 1:
            w.conditional_formatting.add(f"{L}2:{L}{w.max_row}", CellIsRule(operator="lessThan", formula=["45"], font=Font(color="B3261E")))
    wb.save(path)


# ----------------------------------------------------------------------------- HTML
def write_html(res, path):
    tpl = (HERE.parent / "assets" / "acord-alignment.template.html").read_text(encoding="utf-8")
    payload = json.dumps(res, separators=(",", ":")).replace("</", "<\\/")
    title = f"Reference alignment - {res['meta']['region']} ({res['meta']['scope']['slice']})"
    path.write_text(tpl.replace("{{TITLE}}", html.escape(title)).replace("/*__DATA__*/null", payload), encoding="utf-8")


if __name__ == "__main__":
    main()
