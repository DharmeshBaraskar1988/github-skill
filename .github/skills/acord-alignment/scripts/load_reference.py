#!/usr/bin/env python3
"""
Normalise a reference model into reference.json for alignment.

Supported sources (auto-detected by extension / content):
  *.xlsx / *.xlsm   tabular export of your licensed ACORD model (layout in references/reference-model-format.md)
  *.xsd             ACORD XML schema (complexType -> entity, element/attribute -> attribute, simpleType enums -> code lists)
  *.json            already-normalised reference.json, OR a canonical-model.json produced by the canonical-model
                    skill (used as the *baseline* for the next region, e.g. EU canonical for UK)

Usage:
  python load_reference.py --source acord-model.xlsx --name ACORD --version "<your licence version>" --out api-catalog/reference/acord.reference.json
  python load_reference.py --source api-catalog/canonical/EU/canonical-model.json --out api-catalog/reference/eu-baseline.reference.json
  python load_reference.py --sample api-catalog/reference/reference-sample.xlsx    # writes an ILLUSTRATIVE (non-ACORD) sample workbook

The ACORD standards are licensed by ACORD. This kit ships no ACORD content - export the model you are licensed
to use. The sample workbook uses invented generic insurance names so the pipeline can be tested end to end.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

HEADER_ALIASES = {
    "entity": ["entity", "entity name", "class", "object", "aggregate", "complex type", "component"],
    "attribute": ["attribute", "attribute name", "property", "element", "field"],
    "type": ["type", "data type", "datatype", "attribute type"],
    "format": ["format"],
    "required": ["required", "mandatory", "min occurs", "minoccurs", "optionality"],
    "collection": ["collection", "multiple", "max occurs", "maxoccurs", "cardinality", "is list", "repeating"],
    "description": ["description", "definition", "documentation"],
    "subjectArea": ["subject area", "subjectarea", "domain", "area", "package", "model area"],
    "parent": ["parent", "parent entity", "extends", "base", "supertype"],
    "refEntity": ["reference to", "references", "ref entity", "target entity", "related entity"],
    "codeList": ["code list", "codelist", "values", "enumeration", "allowed values"],
    "id": ["id", "acord id", "identifier", "uid", "reference id", "model id"],
    "synonyms": ["synonyms", "aliases", "alias", "also known as"],
}
PRIMS = {
    "string": ("string", ""), "text": ("string", ""), "char": ("string", ""), "varchar": ("string", ""),
    "boolean": ("boolean", ""), "bool": ("boolean", ""), "indicator": ("boolean", ""),
    "integer": ("integer", "int32"), "int": ("integer", "int32"), "long": ("integer", "int64"), "count": ("integer", "int32"),
    "decimal": ("number", "double"), "number": ("number", "double"), "double": ("number", "double"), "float": ("number", "float"),
    "percent": ("number", "double"), "amount": ("number", "double"),
    "date": ("string", "date"), "datetime": ("string", "date-time"), "date-time": ("string", "date-time"),
    "timestamp": ("string", "date-time"), "time": ("string", "time"), "duration": ("string", "duration"),
    "uuid": ("string", "uuid"), "guid": ("string", "uuid"), "uri": ("string", "uri"), "url": ("string", "uri"),
    "code": ("string", ""), "identifier": ("string", ""), "id": ("string", ""), "binary": ("string", "binary"),
    "base64binary": ("string", "byte"), "anyuri": ("string", "uri"), "normalizedstring": ("string", ""),
    "token": ("string", ""), "gyear": ("string", ""), "positiveinteger": ("integer", "int32"),
    "nonnegativeinteger": ("integer", "int32"), "short": ("integer", "int32"),
}


def norm(s) -> str:
    return re.sub(r"\s+", " ", str(s if s is not None else "")).strip()


def truthy(s) -> bool:
    v = norm(s).lower()
    return v in ("y", "yes", "true", "1", "required", "mandatory", "m", "x")


def is_collection(s) -> bool:
    v = norm(s).lower()
    return v in ("y", "yes", "true", "list", "many", "*", "unbounded", "n", "0..*", "1..*", "array") or (v.isdigit() and int(v) > 1)


def map_type(t: str):
    """-> (type, format, refEntity). Unknown non-primitive names become references to entities/code lists."""
    raw = norm(t)
    if not raw:
        return "string", "", None
    key = raw.split(":")[-1].lower()
    key = re.sub(r"_type$", "", key)
    if key in PRIMS:
        typ, fmt = PRIMS[key]
        return typ, fmt, None
    if key.endswith(("datetime", "timestamp")):
        return "string", "date-time", None
    if key.endswith("date"):
        return "string", "date", None
    if key.endswith("indicator") or key.endswith("flag"):
        return "boolean", "", None
    if key.endswith(("text", "name", "identifier", "number")):
        return "string", "", None
    return "object", "", re.sub(r"_Type$", "", raw.split(":")[-1])


def entity_skeleton(name, **kw):
    return {"name": name, "id": kw.get("id", ""), "subjectArea": kw.get("subjectArea", ""),
            "description": kw.get("description", ""), "parent": kw.get("parent"), "kind": "object",
            "codeValues": [], "synonyms": kw.get("synonyms", []), "attributes": []}


# ----------------------------------------------------------------------------- Excel
def load_xlsx(path: Path, sheet: str | None) -> list:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = [wb[sheet]] if sheet else wb.worksheets
    ents: dict = {}
    for ws in sheets:
        rows = list(ws.iter_rows(values_only=True))
        idx, hdr = {}, None
        for i, row in enumerate(rows[:25]):
            cand = {}
            for c, cell in enumerate(row):
                h = norm(cell).lower()
                for k, al in HEADER_ALIASES.items():
                    if h in al and k not in cand:
                        cand[k] = c
            if "entity" in cand:
                idx, hdr = cand, i
                break
        if hdr is None:
            continue
        for row in rows[hdr + 1:]:
            g = lambda k: row[idx[k]] if k in idx and idx[k] < len(row) else None  # noqa: E731
            en = norm(g("entity"))
            if not en:
                continue
            e = ents.setdefault(en, entity_skeleton(en))
            an = norm(g("attribute"))
            for k in ("id", "subjectArea", "parent"):
                if not an and norm(g(k)):
                    e[k] = norm(g(k))
            if not e["subjectArea"] and norm(g("subjectArea")):
                e["subjectArea"] = norm(g("subjectArea"))
            if not an:
                e["description"] = e["description"] or norm(g("description"))
                if norm(g("synonyms")):
                    e["synonyms"] += [x.strip() for x in re.split(r"[,;|]", norm(g("synonyms"))) if x.strip()]
                if norm(g("codeList")):
                    e["kind"] = "enum"
                    e["codeValues"] = [x.strip() for x in re.split(r"[,;|\n]", norm(g("codeList"))) if x.strip()]
                continue
            typ, fmt, ref = map_type(g("type"))
            ref = norm(g("refEntity")) or ref
            if norm(g("format")):
                fmt = norm(g("format"))
            a = {"name": an, "type": "object" if ref else typ, "format": fmt, "refEntity": ref,
                 "isCollection": is_collection(g("collection")), "required": truthy(g("required")),
                 "description": norm(g("description")), "codeValues": [], "id": norm(g("id")),
                 "synonyms": [x.strip() for x in re.split(r"[,;|]", norm(g("synonyms"))) if x.strip()]}
            if norm(g("codeList")):
                a["codeValues"] = [x.strip() for x in re.split(r"[,;|\n]", norm(g("codeList"))) if x.strip()]
            e["attributes"].append(a)
    return list(ents.values())


# ----------------------------------------------------------------------------- XSD
XS = "{http://www.w3.org/2001/XMLSchema}"


def load_xsd(path: Path) -> list:
    """Covers the common XSD shapes (named complexTypes with sequence/choice/all of elements, extension bases,
    attributes, named simpleType enumerations). Element refs resolve via global elements."""
    tree = ET.parse(path)
    root = tree.getroot()
    global_elems = {e.get("name"): e.get("type") for e in root.findall(f"{XS}element") if e.get("name")}
    enums = {}
    for st in root.iter(f"{XS}simpleType"):
        n = st.get("name")
        vals = [en.get("value") for en in st.iter(f"{XS}enumeration")]
        if n and vals:
            enums[re.sub(r"_Type$", "", n)] = vals

    def doc(el):
        d = el.find(f"{XS}annotation/{XS}documentation")
        return norm(d.text) if d is not None and d.text else ""

    ents = []
    for ct in root.iter(f"{XS}complexType"):
        n = ct.get("name")
        if not n:
            continue
        name = re.sub(r"_Type$", "", n)
        e = entity_skeleton(name, description=doc(ct))
        ext = ct.find(f".//{XS}extension")
        if ext is not None and ext.get("base") and not ext.get("base").startswith(("xs:", "xsd:")):
            e["parent"] = re.sub(r"_Type$", "", ext.get("base").split(":")[-1])
        for el in ct.iter(f"{XS}element"):
            en = el.get("name") or (el.get("ref") or "").split(":")[-1]
            if not en:
                continue
            t = el.get("type") or global_elems.get(en, "") or ""
            typ, fmt, ref = map_type(t) if t else ("object", "", en)
            ref_clean = re.sub(r"_Type$", "", ref) if ref else None
            code = enums.get(ref_clean) if ref_clean else None
            e["attributes"].append({
                "name": en, "type": "string" if code else ("object" if ref else typ), "format": fmt,
                "refEntity": None if code else ref_clean, "isCollection": el.get("maxOccurs", "1") not in ("1", "0"),
                "required": el.get("minOccurs", "1") != "0", "description": doc(el), "codeValues": code or [],
                "id": "", "synonyms": []})
        for at in ct.iter(f"{XS}attribute"):
            an = at.get("name")
            if not an:
                continue
            typ, fmt, _ = map_type(at.get("type", "xs:string"))
            e["attributes"].append({"name": an, "type": typ, "format": fmt, "refEntity": None, "isCollection": False,
                                    "required": at.get("use") == "required", "description": doc(at),
                                    "codeValues": enums.get(re.sub(r"_Type$", "", at.get("type", "").split(":")[-1]), []),
                                    "id": "", "synonyms": []})
        ents.append(e)
    for en, vals in enums.items():
        e = entity_skeleton(en)
        e["kind"], e["codeValues"] = "enum", vals
        ents.append(e)
    return ents


# ----------------------------------------------------------------------------- OpenAPI / JSON Schema (YAML or JSON)
def schemas_to_entities(schemas: dict) -> list:
    """OpenAPI components.schemas / JSON Schema $defs|definitions -> entities (ACORD publishes some standards this way)."""
    def ref_name(x):
        return x.split("/")[-1] if isinstance(x, str) else None

    def prop_attr(pn, ps, required):
        ps = ps or {}
        coll = ps.get("type") == "array"
        item = ps.get("items", {}) if coll else ps
        ref = ref_name(item.get("$ref"))
        if not ref and item.get("allOf"):
            ref = next((ref_name(x.get("$ref")) for x in item["allOf"] if isinstance(x, dict) and x.get("$ref")), None)
        typ, fmt = (item.get("type") or ("object" if ref else "string")), item.get("format", "")
        if isinstance(typ, list):
            typ = next((t for t in typ if t != "null"), "string")
        return {"name": pn, "type": "object" if ref else typ, "format": fmt, "refEntity": ref, "isCollection": coll,
                "required": pn in required, "description": norm(ps.get("description") or item.get("description")),
                "codeValues": [str(v) for v in item.get("enum", [])] if not ref else [], "id": "", "synonyms": []}
    ents = []
    for name, sc in (schemas or {}).items():
        sc = sc or {}
        e = entity_skeleton(name, description=norm(sc.get("description") or sc.get("title")))
        e["subjectArea"] = norm(sc.get("x-subject-area") or sc.get("x-acord-subject-area") or "")
        if sc.get("enum"):
            e["kind"], e["codeValues"] = "enum", [str(v) for v in sc["enum"]]
            ents.append(e)
            continue
        parts = sc.get("allOf") or [sc]
        required = set(sc.get("required") or [])
        for part in parts:
            if not isinstance(part, dict):
                continue
            if part.get("$ref") and part is not sc:
                e["parent"] = ref_name(part["$ref"])
                continue
            required |= set(part.get("required") or [])
            for pn, ps in (part.get("properties") or {}).items():
                e["attributes"].append(prop_attr(pn, ps, required))
        ents.append(e)
    return ents


def load_structured(doc: dict, path: Path, name: str, version: str) -> dict:
    if doc.get("kind") in ("reference", "acord") and "entities" in doc:
        return doc
    if doc.get("modelType") == "canonical":
        return load_json(path)
    if "openapi" in doc or "swagger" in doc:
        schemas = (doc.get("components") or {}).get("schemas") or doc.get("definitions") or {}
        return {"kind": "reference", "source": "acord", "name": name, "version": version or str((doc.get("info") or {}).get("version", "")),
                "entities": schemas_to_entities(schemas)}
    if "$defs" in doc or "definitions" in doc or "$schema" in doc:
        schemas = dict(doc.get("$defs") or doc.get("definitions") or {})
        if doc.get("properties") and doc.get("title"):
            schemas.setdefault(re.sub(r"\W", "", doc["title"]), doc)
        return {"kind": "reference", "source": "acord", "name": name, "version": version, "entities": schemas_to_entities(schemas)}
    sys.exit(f"{path}: not a reference model, canonical model, OpenAPI document or JSON Schema")


# ----------------------------------------------------------------------------- JSON / canonical baseline
def load_json(path: Path) -> dict:
    doc = json.loads(path.read_text(encoding="utf-8"))
    if doc.get("kind") in ("reference", "acord") and "entities" in doc:
        return doc
    if doc.get("modelType") == "canonical":
        ents = []
        for ce in doc["entities"]:
            e = entity_skeleton(ce["name"], subjectArea=ce.get("domain", ""), description=ce.get("description", ""))
            e["kind"] = ce.get("kind", "object")
            e["codeValues"] = ce.get("enumValues") or []
            e["id"] = f"{doc.get('region')}:{ce['name']}"
            e["referenceEntity"] = ce.get("reference", {}).get("entity") if ce.get("reference") else None
            for a in ce.get("attributes", []):
                e["attributes"].append({"name": a["name"], "type": a["type"], "format": a.get("format", ""),
                                        "refEntity": a.get("refEntity"), "isCollection": a.get("isCollection", False),
                                        "required": a.get("required", False), "description": a.get("description", ""),
                                        "codeValues": a.get("enum") or [], "id": "", "synonyms": a.get("aliases", [])})
            ents.append(e)
        return {"kind": "reference", "source": "baseline", "name": f"{doc.get('region')} canonical",
                "version": doc.get("version", ""), "region": doc.get("region"), "entities": ents}
    return load_structured(doc, path, "ACORD", "")


def flatten_inheritance(ents: list) -> list:
    by = {e["name"]: e for e in ents}

    def inherited(e, depth=0):
        if not e.get("parent") or depth > 10 or e["parent"] not in by:
            return []
        p = by[e["parent"]]
        return inherited(p, depth + 1) + p["attributes"]
    for e in ents:
        own = {a["name"] for a in e["attributes"]}
        e["inheritedAttributes"] = [dict(a, inherited=True) for a in inherited(e) if a["name"] not in own]
    return ents


def write_sample(path: Path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws.append(["ILLUSTRATIVE SAMPLE - invented generic insurance model for testing the kit. NOT ACORD content. "
               "Replace with an export of your licensed ACORD model."])
    ws["A1"].font = Font(bold=True, color="B3261E")
    ws.append([])
    ws.append(["Entity", "Attribute", "Type", "Required", "Collection", "Reference To", "Code List", "Subject Area", "Description"])
    for c in ws[3]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3A5F")

    rows = [
        ("MonetaryAmount", "", "", "", "", "", "", "Common", "An amount of money in a currency."),
        ("MonetaryAmount", "amount", "decimal", "Y", "", "", "", "", "Numeric value."),
        ("MonetaryAmount", "currencyCode", "code", "Y", "", "", "", "", "ISO 4217 currency code."),
        ("Address", "", "", "", "", "", "", "Common", "Postal address."),
        ("Address", "addressLine1", "string", "Y", "", "", "", "", "First address line."),
        ("Address", "addressLine2", "string", "", "", "", "", "", "Second address line."),
        ("Address", "city", "string", "Y", "", "", "", "", "City or town."),
        ("Address", "postalCode", "string", "", "", "", "", "", "Postal / ZIP code."),
        ("Address", "stateProvince", "string", "", "", "", "", "", "State or province."),
        ("Address", "countryCode", "code", "Y", "", "", "", "", "ISO 3166 country code."),
        ("Party", "", "", "", "", "", "", "Party", "A person or organisation with a role in insurance."),
        ("Party", "partyIdentifier", "identifier", "Y", "", "", "", "", "Unique party id."),
        ("Party", "givenName", "string", "", "", "", "", "", "First name."),
        ("Party", "surname", "string", "", "", "", "", "", "Family name."),
        ("Party", "birthDate", "date", "", "", "", "", "", "Date of birth."),
        ("Party", "emailAddress", "string", "", "", "", "", "", "E-mail address."),
        ("Party", "address", "", "", "", "Address", "", "", "Main postal address."),
        ("ClaimParty", "", "", "", "", "", "", "Claim", "A party involved in a claim (claimant, witness, ...)."),
        ("ClaimParty", "claimPartyIdentifier", "identifier", "Y", "", "", "", "", "Unique id of the claim party."),
        ("ClaimParty", "roleCode", "code", "Y", "", "", "Claimant, Witness, ThirdParty", "", "Role in the claim."),
        ("ClaimParty", "givenName", "string", "", "", "", "", "", "First name."),
        ("ClaimParty", "surname", "string", "", "", "", "", "", "Family name."),
        ("ClaimParty", "emailAddress", "string", "", "", "", "", "", "E-mail address."),
        ("ClaimParty", "address", "", "", "", "Address", "", "", "Postal address."),
        ("ClaimStatusCode", "", "", "", "", "", "Open, Submitted, UnderReview, Approved, Denied, Closed, Reopened", "Claim", "Claim lifecycle status."),
        ("Claim", "", "", "", "", "", "", "Claim", "A demand for payment under a policy following a loss."),
        ("Claim", "claimIdentifier", "uuid", "Y", "", "", "", "", "Unique claim id."),
        ("Claim", "claimNumber", "string", "Y", "", "", "", "", "Business claim number."),
        ("Claim", "policyNumber", "string", "Y", "", "", "", "", "Number of the policy claimed against."),
        ("Claim", "claimStatusCode", "", "", "", "ClaimStatusCode", "", "", "Lifecycle status."),
        ("Claim", "lossDate", "date", "Y", "", "", "", "", "Date of loss."),
        ("Claim", "reportedDate", "date", "", "", "", "", "", "Date the loss was reported."),
        ("Claim", "lossDescription", "string", "", "", "", "", "", "Narrative of the loss."),
        ("Claim", "lossLocation", "", "", "", "Address", "", "", "Where the loss happened."),
        ("Claim", "reserveAmount", "", "", "", "MonetaryAmount", "", "", "Current reserve."),
        ("Claim", "claimParties", "", "", "Y", "ClaimParty", "", "", "Parties involved."),
        ("Document", "", "", "", "", "", "", "Common", "A document attached to a business object."),
        ("Document", "documentIdentifier", "uuid", "Y", "", "", "", "", "Unique document id."),
        ("Document", "fileName", "string", "Y", "", "", "", "", "Original file name."),
        ("Document", "documentTypeCode", "code", "", "", "", "", "", "Kind of document."),
        ("Document", "fileSize", "long", "", "", "", "", "", "Size in bytes."),
        ("Document", "mimeType", "string", "", "", "", "", "", "Media type."),
        ("PolicyStatusCode", "", "", "", "", "", "Quoted, InForce, Lapsed, Cancelled, Expired", "Policy", "Policy status."),
        ("Coverage", "", "", "", "", "", "", "Policy", "A coverage provided by a policy."),
        ("Coverage", "coverageCode", "code", "Y", "", "", "", "", "Coverage code."),
        ("Coverage", "coverageName", "string", "", "", "", "", "", "Coverage name."),
        ("Coverage", "limitAmount", "", "", "", "MonetaryAmount", "", "", "Limit."),
        ("Coverage", "deductibleAmount", "", "", "", "MonetaryAmount", "", "", "Deductible."),
        ("Policy", "", "", "", "", "", "", "Policy", "An insurance contract."),
        ("Policy", "policyNumber", "string", "Y", "", "", "", "", "Policy number."),
        ("Policy", "productCode", "code", "Y", "", "", "", "", "Product code."),
        ("Policy", "policyStatusCode", "", "", "", "PolicyStatusCode", "", "", "Status."),
        ("Policy", "effectiveDate", "date", "Y", "", "", "", "", "Start of cover."),
        ("Policy", "expirationDate", "date", "Y", "", "", "", "", "End of cover."),
        ("Policy", "premiumAmount", "", "", "", "MonetaryAmount", "", "", "Total premium."),
        ("Policy", "policyHolder", "", "", "", "Party", "", "", "Policy holder."),
        ("Policy", "coverages", "", "", "Y", "Coverage", "", "", "Coverages."),
        ("Quote", "", "", "", "", "", "", "Quote", "A priced offer for a policy."),
        ("Quote", "quoteIdentifier", "string", "Y", "", "", "", "", "Quote reference."),
        ("Quote", "productCode", "code", "Y", "", "", "", "", "Product code."),
        ("Quote", "premiumAmount", "", "", "", "MonetaryAmount", "", "", "Quoted premium."),
        ("Quote", "expirationDate", "date", "", "", "", "", "", "Quote valid until."),
        ("Quote", "proposedHolder", "", "", "", "Party", "", "", "Prospective policy holder."),
        ("Quote", "effectiveDate", "date", "", "", "", "", "", "Requested start date."),
    ]
    for r in rows:
        ws.append(list(r))
    for col, w in zip("ABCDEFGHI", (20, 24, 12, 10, 11, 18, 40, 14, 50)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--source")
    ap.add_argument("--sheet")
    ap.add_argument("--name", default="ACORD")
    ap.add_argument("--version", default="")
    ap.add_argument("--out")
    ap.add_argument("--sample", help="write the illustrative sample workbook to this path and exit")
    a = ap.parse_args()
    if a.sample:
        Path(a.sample).parent.mkdir(parents=True, exist_ok=True)
        write_sample(Path(a.sample))
        print(f"Illustrative (non-ACORD) sample written: {a.sample}")
        return
    if not a.source or not a.out:
        ap.error("--source and --out are required")
    src = Path(a.source)
    if src.suffix.lower() in (".xlsx", ".xlsm"):
        doc = {"kind": "reference", "source": "acord", "name": a.name, "version": a.version,
               "entities": flatten_inheritance(load_xlsx(src, a.sheet))}
    elif src.suffix.lower() == ".xsd":
        doc = {"kind": "reference", "source": "acord", "name": a.name, "version": a.version,
               "entities": flatten_inheritance(load_xsd(src))}
    elif src.suffix.lower() == ".json":
        doc = load_json(src)
        doc["name"], doc["version"] = (a.name if doc.get("source") == "acord" else doc["name"]), (a.version or doc.get("version", ""))
        doc["entities"] = flatten_inheritance(doc["entities"])
    elif src.suffix.lower() in (".yaml", ".yml"):
        import yaml
        doc = load_structured(yaml.safe_load(src.read_text(encoding="utf-8")) or {}, src, a.name, a.version)
        doc["entities"] = flatten_inheritance(doc["entities"])
    elif src.is_dir():
        # a folder of ACORD files (e.g. one YAML/JSON/XSD per subject area) -> merged
        import yaml
        ents, seen = [], set()
        for f in sorted(src.rglob("*")):
            if f.suffix.lower() in (".yaml", ".yml", ".json"):
                d = yaml.safe_load(f.read_text(encoding="utf-8")) if f.suffix != ".json" else json.loads(f.read_text(encoding="utf-8"))
                part = load_structured(d or {}, f, a.name, a.version)["entities"]
            elif f.suffix.lower() == ".xsd":
                part = load_xsd(f)
            elif f.suffix.lower() in (".xlsx", ".xlsm"):
                part = load_xlsx(f, a.sheet)
            else:
                continue
            for e in part:
                if e["name"] not in seen:
                    seen.add(e["name"])
                    ents.append(e)
        doc = {"kind": "reference", "source": "acord", "name": a.name, "version": a.version, "entities": flatten_inheritance(ents)}
    else:
        sys.exit(f"Unsupported reference source: {src}")
    doc["sourceFile"] = src.name
    for e in doc["entities"]:
        if e["kind"] != "enum" and not e["attributes"] and e.get("codeValues"):
            e["kind"] = "enum"
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    Path(a.out).write_text(json.dumps(doc, indent=1), encoding="utf-8")
    print(json.dumps({"output": a.out, "source": doc["source"], "name": doc["name"], "version": doc.get("version", ""),
                      "entities": len(doc["entities"]),
                      "codeLists": sum(1 for e in doc["entities"] if e["kind"] == "enum"),
                      "attributes": sum(len(e["attributes"]) for e in doc["entities"])}, indent=2))


if __name__ == "__main__":
    main()
