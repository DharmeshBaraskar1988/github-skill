#!/usr/bin/env python3
"""
Build the approved canonical model of a region from a reviewed alignment.

Only rows a human approved (approvals.yaml, written by import_review.py) are used. Pending or changed rows
make the result a 'draft' (validation C01 fails) unless canonical.allowPending is set, in which case they are
left out and listed.

Inputs  (region config, see acord-alignment/templates/region.config.template.yaml)
  alignment/<REGION>/alignment.json + approvals.yaml
  reference.json (ACORD export) and optional baseline canonical-model.json (e.g. EU when building UK)
  regional-view-data.json (for operation contracts and schema lineage)

Outputs (canonical/<REGION>/)
  canonical-model.json            canonical entities, attributes, relations, endpoints, lineage (machine readable;
                                  also the baseline input for the next region)
  canonical-openapi.yaml / .json  ONE OpenAPI 3.0 spec: canonical endpoints + canonical schemas
  source-to-canonical-mapping.json  app schema.attribute -> canonical entity.attribute (+ reference attribute)
  canonical-model.xlsx            Entities, Attributes, Relations, Endpoints, Source mapping, Changes, Excluded
  canonical-model.md, CHANGELOG.md
  releases/<version>/             frozen copy when --release is given

Usage
  python build_canonical.py --config api-catalog/regions/EU.yaml
  python build_canonical.py --config api-catalog/regions/EU.yaml --release        # freeze canonical.version
"""
from __future__ import annotations

import argparse
import copy
import datetime as dt
import json
import os
import re
import shutil
import sys
from collections import defaultdict
from pathlib import Path

import yaml

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent.parent / "acord-alignment" / "scripts"))
sys.path.insert(0, str(HERE))
from align_core import Vocab, proposed_path, review_status, base_type  # noqa: E402
from load_reference import load_json, flatten_inheritance  # noqa: E402

PRIM_FORMATS = {"int32", "int64", "double", "float", "date", "date-time", "time", "uuid", "uri", "email", "byte",
                "binary", "duration"}


class NoAlias(yaml.SafeDumper):
    def ignore_aliases(self, data):
        return True


def rp(base, p):
    if not p:
        return None
    p = Path(p)
    return p if p.is_absolute() else (base / p).resolve()


def pascal(s: str) -> str:
    return "".join(w[:1].upper() + w[1:] for w in re.split(r"[^A-Za-z0-9]+", s or "") if w)


def parse_type(t: str, fmt: str = ""):
    """'string(date-time)' / 'integer(int32)[]' / 'map<string,X>' -> (type, format, isArray, mapValue)"""
    t = (t or "").strip()
    arr = t.endswith("[]")
    if arr:
        t = t[:-2]
    if t.startswith("map<"):
        return "object", "", arr, t[4:-1].split(",", 1)[1] if "," in t else "string"
    m = re.match(r"^(\w+)\((.+)\)$", t)
    if m:
        return m.group(1), m.group(2), arr, None
    return t or "string", fmt or "", arr, None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", required=True)
    ap.add_argument("--alignment-dir")
    ap.add_argument("--release", action="store_true")
    a = ap.parse_args()
    cfg = yaml.safe_load(open(a.config, encoding="utf-8")) or {}
    base = Path(a.config).resolve().parent
    ccfg = cfg.get("canonical", {}) or {}
    region = cfg["region"]
    al_dir = rp(base, a.alignment_dir or (cfg.get("alignment") or {}).get("outputDir") or f"../alignment/{region}")
    out = rp(base, ccfg.get("outputDir") or f"../canonical/{region}")
    out.mkdir(parents=True, exist_ok=True)
    al = json.loads((al_dir / "alignment.json").read_text(encoding="utf-8"))
    if al["meta"]["scope"].get("apps") or al["meta"]["scope"].get("domains") or al["meta"]["scope"].get("entities"):
        sys.exit("The canonical model must be built from a whole-region alignment (no --apps/--domains/--entities slice).")
    approvals = yaml.safe_load((al_dir / "approvals.yaml").read_text(encoding="utf-8")) if (al_dir / "approvals.yaml").exists() else {}
    approvals = approvals or {}
    data = json.loads((al_dir / al["meta"]["regionalViewData"]).resolve().read_text(encoding="utf-8"))
    refs = {}
    baseline = None
    base_p = rp(base, cfg.get("baseline"))
    if base_p:
        baseline = json.loads(base_p.read_text(encoding="utf-8"))
        refs["baseline"] = flatten_inheritance(load_json(base_p)["entities"])
    ref_p = rp(base, (cfg.get("reference") or {}).get("json"))
    if ref_p and ref_p.exists():
        refdoc = json.loads(ref_p.read_text(encoding="utf-8"))
        refs["acord"] = flatten_inheritance(refdoc["entities"])
    ref_index = {(src, e["name"]): e for src, ents in refs.items() for e in ents}
    prev_p = out / "canonical-model.json"
    previous = json.loads(prev_p.read_text(encoding="utf-8")) if prev_p.exists() else None

    model = Builder(al, approvals, data, ref_index, baseline, ccfg, region, cfg).build()
    model["generatedAt"] = dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")
    rel = lambda p: os.path.relpath(p, out)  # noqa: E731
    model["inputs"] = {"alignment": rel(al_dir / "alignment.json"), "alignmentGeneratedAt": al["meta"]["generatedAt"],
                       "approvals": rel(al_dir / "approvals.yaml"), "baseline": rel(base_p) if base_p else None,
                       "reference": [r for r in al["references"]]}
    model["changes"] = diff_models(previous, model, "previous") if previous else {"note": "first build"}
    if baseline:
        model["changesVsBaseline"] = diff_models(baseline, model, "baseline")
    spec = to_openapi(model, ccfg)
    from sensitive_scan import find_policy_file, load_policy, redact_obj
    pol = load_policy(find_policy_file(out))
    model, _ = redact_obj(model, pol)
    spec, _ = redact_obj(spec, pol)
    write_all(out, model, spec)
    if a.release:
        if model["status"] != "approved":
            sys.exit("Refusing to release a draft (pending items). Finish the review or set allowPending.")
        rel = out / "releases" / model["version"]
        if rel.exists():
            sys.exit(f"Release {model['version']} already exists - bump canonical.version in the region config.")
        rel.mkdir(parents=True)
        for f in ("canonical-model.json", "canonical-model.yaml", "canonical-openapi.yaml", "canonical-openapi.json",
                  "source-to-canonical-mapping.json", "canonical-model.xlsx", "canonical-model.md", "CHANGELOG.md",
                  "canonical-viewer.html"):
            shutil.copy2(out / f, rel / f)
    s = model["summary"]
    print(json.dumps({"output": str(out), "status": model["status"], "version": model["version"], **s,
                      "released": bool(a.release)}, indent=2))


class Builder:
    def __init__(self, al, approvals, data, ref_index, baseline, ccfg, region, cfg):
        self.al, self.ap, self.data, self.ref_index = al, approvals, data, ref_index
        self.baseline, self.ccfg, self.region, self.cfg = baseline, ccfg, region, cfg
        self.vocab = Vocab()
        self.allow_pending = bool(ccfg.get("allowPending"))
        self.problems = []           # blocking, reported by the validator
        self.pending, self.rejected, self.excluded_ops, self.lineage = [], [], [], []
        self.entities = {}           # canonical name -> entity
        self.row_canon = {}          # regional entity name -> canonical name (approved only)
        self.ref_to_canon = {}       # (source, reference entity) -> canonical name

    # ------------------------------------------------------------------ helpers
    def ent_approval(self, row):
        a = (self.ap.get("entities") or {}).get(row["id"])
        st = review_status(a, row["basis"])
        return a or {}, st

    def effective(self, row, a):
        dec = a.get("decision")
        if dec == "approve":
            dec = row["recommendation"]
        ref = a.get("reference") or (row["match"]["reference"] if row["match"] else None)
        src = a.get("source") or row.get("referenceSource")
        if dec == "custom":
            ref, src = None, None
        cand = None
        if ref:
            cand = next((c for c in row["candidates"] if c["reference"] == ref and (not src or c["source"] == src)), None)
            if not cand and row["match"] and row["match"]["reference"] == ref:
                cand = row["match"]
            if cand:
                src = cand["source"]
        if dec in ("use-as-is", "extend") and not cand:
            self.problems.append({"code": "C02", "message": f"{row['name']}: decision '{dec}' but reference '{ref}' is not a candidate",
                                  "fix": "Pick a reference from the candidates in review, or decide custom"})
            dec = "custom"
        return dec, cand, src

    def new_entity(self, name, kind, row, approach, cand, src, a):
        ref_e = self.ref_index.get((src, cand["reference"])) if cand else None
        return {"name": name, "kind": kind, "domain": row["domain"], "domains": [row["domain"]] if row["domain"] else [],
                "description": (ref_e or {}).get("description") or row["description"] or "",
                "approach": approach, "reference": ({"source": src, "entity": cand["reference"],
                                                     "subjectArea": (ref_e or {}).get("subjectArea", "")} if cand else None),
                "introducedIn": self.region, "status": "approved",
                "attributes": {}, "enumValues": [], "regionalEntities": [], "sourceSchemas": [],
                "approvedBy": sorted({a.get("reviewer", "")} - {""}), "comments": [a["comment"]] if a.get("comment") else []}

    # ------------------------------------------------------------------ build
    def build(self):
        # 0. baseline entities are kept as they are
        if self.baseline:
            for be in self.baseline["entities"]:
                e = copy.deepcopy(be)
                e["attributes"] = {x["name"]: x for x in be["attributes"]}
                e["status"] = "baseline"
                e.setdefault("regionalEntities", [])
                e.setdefault("sourceSchemas", [])
                self.entities[e["name"]] = e
                self.ref_to_canon[("baseline", e["name"])] = e["name"]
                if e.get("reference"):
                    self.ref_to_canon.setdefault((e["reference"]["source"], e["reference"]["entity"]), e["name"])
        rows = self.al["entities"]
        decided = []
        for row in rows:
            a, st = self.ent_approval(row)
            if st in ("pending", "changed"):
                self.pending.append({"entity": row["name"], "state": st})
                continue
            if st == "rejected":
                self.rejected.append({"entity": row["name"], "reviewer": a.get("reviewer"), "comment": a.get("comment", "")})
                for at in row["attributes"]:
                    for s in at["sources"]:
                        self.lineage.append(self.lin(s, row, at, None, None, "entity rejected"))
                continue
            dec, cand, src = self.effective(row, a)
            if src == "baseline":
                canon = a.get("canonicalName") or cand["reference"]
            else:
                canon = a.get("canonicalName") or (self.ref_to_canon.get((src, cand["reference"])) if cand else None) \
                    or (cand["reference"] if cand else row["name"])
            self.row_canon[row["name"]] = canon
            if cand:
                self.ref_to_canon.setdefault((src, cand["reference"]), canon)
            decided.append((row, a, dec, cand, src, canon))
        # 1. entities + attributes (second loop so references between entities can be resolved)
        for row, a, dec, cand, src, canon in decided:
            approach = {"use-as-is": "reference-as-is", "extend": "reference-extended", "custom": "custom"}[dec]
            if src == "baseline":
                approach = {"reference-as-is": "baseline-reused", "reference-extended": "baseline-extended"}.get(approach, approach)
            e = self.entities.get(canon)
            if not e:
                e = self.entities[canon] = self.new_entity(canon, row["kind"], row, approach, cand, src, a)
            else:
                if row["domain"] and row["domain"] not in e.setdefault("domains", []):
                    e["domains"].append(row["domain"])
                e.setdefault("approvedBy", [])
                if a.get("reviewer") and a["reviewer"] not in e["approvedBy"]:
                    e["approvedBy"].append(a["reviewer"])
                if e["status"] == "baseline" and dec == "extend":
                    e["status"] = "baseline-extended"
            e["regionalEntities"].append({"entity": row["name"], "decision": dec, "alignmentPercent": row["alignmentPercent"],
                                          "apps": row["apps"]})
            for m in row["members"]:
                for sc in m["schemas"]:
                    e["sourceSchemas"].append(f"{m['appId']}:{sc}")
            if row["kind"] == "enum":
                self.enum_values(e, row, dec, cand, src)
                continue
            self.attributes(e, row, dec, cand, src)
            if dec in ("use-as-is", "extend") and self.ccfg.get("includeUnmatchedReferenceAttributes") and cand:
                ref_e = self.ref_index.get((src, cand["reference"]))
                for ra in (ref_e or {}).get("attributes", []) + (ref_e or {}).get("inheritedAttributes", []):
                    if ra["name"] not in e["attributes"]:
                        e["attributes"][ra["name"]] = self.attr_from_ref(ra, src, cand["reference"], reference_only=True)
        # 2. resolve attribute references to canonical names
        for e in self.entities.values():
            for at in e["attributes"].values():
                if at.get("_refRegional"):
                    tgt = self.row_canon.get(at["_refRegional"])
                    if not tgt:
                        self.problems.append({"code": "C02", "message": f"{e['name']}.{at['name']} refers to '{at['_refRegional']}' which is not approved (pending/rejected)",
                                              "fix": "Approve that entity, or exclude/rename this attribute in review"})
                    at["refEntity"] = tgt or at["_refRegional"]
                if at.get("_refReference"):
                    src, rn = at["_refReference"]
                    tgt = self.ref_to_canon.get((src, rn))
                    if not tgt:
                        ref_e = self.ref_index.get((src, rn))
                        if ref_e and ref_e.get("kind") == "enum":
                            tgt = rn
                            if rn not in self.entities:
                                self.entities[rn] = {"name": rn, "kind": "enum", "domain": e["domain"], "domains": e.get("domains", []),
                                                     "description": ref_e.get("description", ""), "approach": "reference-as-is",
                                                     "reference": {"source": src, "entity": rn, "subjectArea": ref_e.get("subjectArea", "")},
                                                     "introducedIn": self.region, "status": "approved", "attributes": {},
                                                     "enumValues": ref_e.get("codeValues", []), "regionalEntities": [],
                                                     "sourceSchemas": [], "approvedBy": [], "comments": ["pulled in as code list of an approved attribute"]}
                        elif at.get("_fallbackRegional") and self.row_canon.get(at["_fallbackRegional"]):
                            tgt = self.row_canon[at["_fallbackRegional"]]
                        else:
                            self.problems.append({"code": "C02", "message": f"{e['name']}.{at['name']} uses reference entity '{rn}' that has no approved canonical entity",
                                                  "fix": "Approve the regional entity that maps to it, or exclude this attribute"})
                            tgt = rn
                    at["refEntity"] = tgt
                for k in ("_refRegional", "_refReference", "_fallbackRegional"):
                    at.pop(k, None)
        # 3. endpoints
        endpoints = self.endpoints()
        # 4. relations
        relations = []
        for e in self.entities.values():
            for at in e["attributes"].values():
                if at.get("refEntity") and at["refEntity"] in self.entities:
                    k = self.entities[at["refEntity"]]["kind"]
                    relations.append({"from": e["name"], "to": at["refEntity"], "via": at["name"],
                                      "type": "uses-code-list" if k == "enum" else ("has-many" if at.get("isCollection") else "has-one"),
                                      "cardinality": "0..*" if at.get("isCollection") else ("1" if at.get("required") else "0..1")})
        ents = []
        for e in sorted(self.entities.values(), key=lambda x: x["name"]):
            e2 = dict(e)
            e2["attributes"] = list(e["attributes"].values())
            e2["sourceSchemas"] = sorted(set(e.get("sourceSchemas", [])))
            ents.append(e2)
        status = "approved" if not self.pending or self.allow_pending else "draft"
        reviewers = sorted({v.get("reviewer") for k in ("entities", "attributes", "endpoints")
                            for v in (self.ap.get(k) or {}).values() if v.get("reviewer")})
        refs = [{"source": r["source"], "name": r["name"], "version": r["version"]} for r in self.al["references"]]
        return {
            "modelType": "canonical", "region": self.region, "version": str(self.ccfg.get("version", "1.0.0")),
            "title": self.ccfg.get("title") or f"{self.region} Canonical Insurance API", "status": status,
            "baseline": {"region": self.baseline["region"], "version": self.baseline["version"]} if self.baseline else None,
            "references": refs, "reviewers": reviewers,
            "entities": ents, "relations": relations, "endpoints": endpoints, "lineage": self.lineage,
            "pending": self.pending, "rejected": self.rejected, "excludedEndpoints": self.excluded_ops,
            "problems": self.problems,
            "summary": {"entities": len(ents), "attributes": sum(len(e["attributes"]) for e in ents),
                        "baselineEntities": sum(1 for e in ents if e["status"].startswith("baseline")),
                        "newEntities": sum(1 for e in ents if e["status"] == "approved"),
                        "extensions": sum(1 for e in ents for at in e["attributes"] if at.get("extension")),
                        "endpoints": len(endpoints), "relations": len(relations), "pending": len(self.pending),
                        "rejected": len(self.rejected), "excludedEndpoints": len(self.excluded_ops),
                        "lineageRows": len(self.lineage), "problems": len(self.problems)},
        }

    def lin(self, s, row, at, canon_e, canon_a, note, ref=None, transform=""):
        return {"region": self.region, "appId": s["appId"], "sourceSchemas": s["schemas"], "sourceAttribute": s["attribute"],
                "appEntity": s["entity"], "regionalEntity": row["name"], "regionalAttribute": at["name"],
                "canonicalEntity": canon_e, "canonicalAttribute": canon_a, "referenceAttribute": ref,
                "transformation": transform, "note": note}

    def attr_from_ref(self, ra, src, ref_entity, reference_only=False):
        at = {"name": ra["name"], "type": "object" if ra.get("refEntity") else ra.get("type", "string"),
              "format": ra.get("format", ""), "refEntity": None, "isCollection": bool(ra.get("isCollection")),
              "required": bool(ra.get("required")), "description": ra.get("description", ""),
              "reference": {"source": src, "entity": ref_entity, "attribute": ra["name"]}, "extension": False,
              "enum": ra.get("codeValues") or None, "sources": [], "introducedIn": self.region}
        if ra.get("refEntity"):
            at["_refReference"] = (src, ra["refEntity"])
        if reference_only:
            at["referenceOnly"] = True
        return at

    def attributes(self, e, row, dec, cand, src):
        aap = self.ap.get("attributes") or {}
        matches = {x["name"]: x for x in (cand or {}).get("attributes", [])}
        ref_e = self.ref_index.get((src, cand["reference"])) if cand else None
        ref_attrs = {x["name"]: x for x in ((ref_e or {}).get("attributes", []) + (ref_e or {}).get("inheritedAttributes", []))}
        for at in row["attributes"]:
            ad = aap.get(at["id"]) or {}
            d = ad.get("decision") or "accept"
            mt = matches.get(at["name"]) or {}
            if d == "exclude":
                for s in at["sources"]:
                    self.lineage.append(self.lin(s, row, at, e["name"], None, f"excluded by {ad.get('reviewer', '?')}: {ad.get('comment', '')}"))
                continue
            ref_name = mt.get("reference") if mt.get("status") != "no-match" else None
            if dec == "custom":
                ref_name = None
            if dec == "use-as-is" and not ref_name:
                for s in at["sources"]:
                    self.lineage.append(self.lin(s, row, at, e["name"], None, "dropped: entity adopted reference as-is"))
                continue
            if ref_name and ref_name in ref_attrs:
                cname = (ad.get("canonicalName") if d == "rename" else None) or ref_name
                new = self.attr_from_ref(ref_attrs[ref_name], src, cand["reference"])
                new["name"] = cname
                if not new["description"]:
                    new["description"] = at.get("description", "")
                if at.get("refEntity") and new.get("_refReference"):
                    new["_fallbackRegional"] = at["refEntity"]
            else:
                cname = (ad.get("canonicalName") if d == "rename" else None) or at["name"]
                t, f, arr, _ = parse_type(at["type"], at.get("format", ""))
                new = {"name": cname, "type": "object" if at.get("refEntity") else (t if t not in ("array",) else "string"),
                       "format": f if f in PRIM_FORMATS else (at.get("format") if at.get("format") in PRIM_FORMATS else ""),
                       "refEntity": None, "isCollection": bool(at.get("isCollection")) or arr,
                       "required": bool(at.get("required")), "description": at.get("description", ""),
                       "reference": None, "extension": dec in ("extend",) or e["status"].startswith("baseline"),
                       "enum": None, "sources": [], "introducedIn": self.region}
                if at.get("refEntity"):
                    new["_refRegional"] = at["refEntity"]
            existing = e["attributes"].get(cname)
            if existing:
                t_old = (existing.get("refEntity") or existing.get("_refRegional") or existing.get("type"), existing.get("isCollection"))
                t_new = (new.get("refEntity") or new.get("_refRegional") or new.get("type"), new.get("isCollection"))
                if existing.get("_refReference") or new.get("_refReference"):
                    pass
                elif base_type(str(t_old[0])) != base_type(str(t_new[0])) and not existing.get("refEntity") and not new.get("_refRegional"):
                    self.problems.append({"code": "C04", "message": f"{e['name']}.{cname}: type {t_old} from one source vs {t_new} from {row['name']}",
                                          "fix": "Rename one attribute in review or align the types at source"})
                if not existing.get("description") and new.get("description"):
                    existing["description"] = new["description"]
                tgt = existing
            else:
                e["attributes"][cname] = new
                tgt = new
            transform = []
            if cname != at["name"]:
                transform.append(f"rename {at['name']} -> {cname}")
            if ref_name and ref_attrs.get(ref_name) and not at.get("refEntity"):
                rt = ref_attrs[ref_name]
                if base_type(at["type"]) != rt.get("type") or (rt.get("format") and rt.get("format") != at.get("format")):
                    transform.append(f"type {at['type']} -> {rt.get('type')}{'(' + rt['format'] + ')' if rt.get('format') else ''}")
            for s in at["sources"]:
                ref_label = f"{src}:{cand['reference']}.{ref_name}" if ref_name else None
                tgt.setdefault("sources", []).append(f"{s['appId']}:{'/'.join(s['schemas'])}.{s['attribute']}")
                self.lineage.append(self.lin(s, row, at, e["name"], cname, "mapped" if ref_name else ("extension" if dec == "extend" else "custom"),
                                             ref_label, "; ".join(transform)))

    def enum_values(self, e, row, dec, cand, src):
        ref_e = self.ref_index.get((src, cand["reference"])) if cand else None
        ref_vals = list((ref_e or {}).get("codeValues", []))
        ours = row["enumValues"]
        if dec == "use-as-is":
            vals = ref_vals
        elif dec == "extend":
            norm = {"".join(self.vocab.tokens(v)) for v in ref_vals}
            vals = ref_vals + [v for v in ours if "".join(self.vocab.tokens(v)) not in norm]
            e["extensionValues"] = [v for v in ours if "".join(self.vocab.tokens(v)) not in norm]
        else:
            vals = ours
        for v in vals:
            if v not in e["enumValues"]:
                e["enumValues"].append(v)
        e["valueMapping"] = [{"source": v, "canonical": next((r for r in vals if "".join(self.vocab.tokens(r)) == "".join(self.vocab.tokens(v))), None)}
                             for v in ours]

    # ------------------------------------------------------------------ endpoints
    def endpoints(self):
        oap = self.ap.get("endpoints") or {}
        explicit = (self.ccfg.get("endpointApproval") or "implicit") == "explicit"
        drop_rejected = self.ccfg.get("dropEndpointsWithRejectedEntities", True)
        rejected = {r["entity"] for r in self.rejected}
        pending = {p["entity"] for p in self.pending}
        app_info = {x["id"]: x for x in self.data["apps"]}
        a2r = self.al.get("appEntityToRegional", {})
        ops_by_id = {f"{o['appId']}:{o['operationId']}": o for o in self.data["operations"]}
        canon_names = dict(self.row_canon)
        for r in self.al["entities"]:
            for m in r["members"]:
                if r["name"] in self.row_canon:
                    canon_names.setdefault(m["entity"], self.row_canon[r["name"]])
        groups = {}
        for ep in self.al["endpoints"]:
            a = oap.get(ep["id"]) or {}
            st = review_status(a, ep["basis"]) if a else ("pending" if explicit else "default")
            if a.get("decision") == "exclude":
                self.excluded_ops.append({"id": ep["id"], "reason": f"excluded by {a.get('reviewer')}: {a.get('comment', '')}"})
                continue
            if st in ("pending", "changed"):
                self.pending.append({"endpoint": ep["id"], "state": st})
                continue
            bad = [x for x in ep["entities"] if x in rejected]
            if bad:
                if drop_rejected:
                    self.excluded_ops.append({"id": ep["id"], "reason": f"uses rejected entities {bad}"})
                else:
                    self.problems.append({"code": "C05", "message": f"{ep['id']} uses rejected entities {bad}",
                                          "fix": "Exclude the endpoint in review or approve the entity"})
                continue
            waiting = [x for x in ep["entities"] if x in pending]
            if waiting:
                self.pending.append({"endpoint": ep["id"], "state": f"waiting for entities {waiting}"})
                continue
            o = ops_by_id.get(ep["id"])
            if not o:
                continue
            path = a.get("canonicalPath") or proposed_path(ep["path"], {}, canon_names, self.vocab)
            key = (ep["method"], path)
            g = groups.setdefault(key, {"method": ep["method"], "path": path, "sources": [], "domain": ep["domain"],
                                        "capability": ep["capability"], "summary": ep.get("summary") or ep["capability"],
                                        "operationId": a.get("canonicalOperationId") or "", "contract": None,
                                        "variants": [], "approvedBy": []})
            if a.get("reviewer"):
                g["approvedBy"].append(a["reviewer"])
            if a.get("canonicalOperationId"):
                g["operationId"] = a["canonicalOperationId"]
            g["sources"].append({"id": ep["id"], "path": ep["path"], "operationId": ep["operationId"]})
            contract = self.map_contract(o, app_info.get(o["appId"], {}), a2r)
            if g["contract"] is None:
                g["contract"] = contract
            elif json.dumps(contract, sort_keys=True) != json.dumps(g["contract"], sort_keys=True):
                g["variants"].append({"source": ep["id"], "contract": contract})
        used, out = defaultdict(int), []
        for (m, p), g in sorted(groups.items(), key=lambda kv: (kv[0][1], kv[0][0])):
            oid = g["operationId"] or pascal(g["capability"]) or pascal(m.lower() + " " + " ".join(s for s in p.split("/") if s and not s.startswith("{")))
            used[oid] += 1
            if used[oid] > 1:
                oid = f"{oid}{used[oid]}"
            g["operationId"] = oid
            out.append(g)
        return out

    def map_contract(self, o, app, a2r):
        s2e = app.get("schemaToEntity", {})
        env = {x["schema"]: x for x in app.get("excludedSchemas", []) if x.get("reason") == "envelope"}
        problem = {x["schema"] for x in app.get("excludedSchemas", []) if x.get("reason") == "problem"}

        def canon(schema):
            if not schema:
                return None
            if schema in problem:
                return {"problem": True}
            if schema in env:
                inner = canon(env[schema]["itemsSchema"])
                return {"page": (inner or {}).get("entity"), "itemsProperty": env[schema]["itemsProperty"],
                        "other": env[schema]["otherProperties"]}
            ent = s2e.get(schema)
            reg = a2r.get(f"{o['appId']}:{ent}") if ent else None
            c = self.row_canon.get(reg) if reg else None
            if not c:
                self.problems.append({"code": "C02", "message": f"{o['appId']}:{o['operationId']} uses schema {schema} with no approved canonical entity",
                                      "fix": "Approve the entity or exclude the endpoint"})
                return {"entity": None, "unresolved": schema}
            return {"entity": c}
        req = o.get("request")
        resp = {}
        for code, r in (o.get("responses") or {}).items():
            if not r:
                resp[code] = None
                continue
            m = canon(r.get("schema")) if r.get("schema") else {"primitive": r.get("type")}
            if m is not None:
                m["isArray"] = r.get("isArray", False)
            resp[code] = m
        params = []
        for p in o.get("parameters", []):
            pm = {"name": p["name"], "in": p["in"], "required": p["required"]}
            if p.get("schema"):
                pm["entity"] = (canon(p["schema"]) or {}).get("entity")
            else:
                pm["type"] = p["type"]
                pm["format"] = p.get("format", "")
            pm["isArray"] = p.get("isArray", False)
            params.append(pm)
        rq = None
        if req:
            rq = canon(req.get("schema")) if req.get("schema") else (
                {"inline": req["inline"]} if req.get("inline") else {"primitive": req.get("type")})
            if rq is not None:
                rq["isArray"] = req.get("isArray", False)
                rq["contentType"] = req.get("contentType", "application/json")
        return {"request": rq, "responses": resp, "parameters": params, "security": o.get("security", [])}


# ----------------------------------------------------------------------------- OpenAPI
def attr_schema(at):
    if at.get("refEntity"):
        s = {"$ref": f"#/components/schemas/{at['refEntity']}"}
    else:
        t, f, arr, mapv = parse_type(at.get("type", "string"), at.get("format", ""))
        if t in ("array",):
            t = "string"
        s = {"type": t if t in ("string", "integer", "number", "boolean", "object") else "string"}
        if f and f in PRIM_FORMATS:
            s["format"] = f
        if at.get("enum"):
            s["enum"] = at["enum"]
    if at.get("isCollection"):
        s = {"type": "array", "items": s}
    meta = {}
    if at.get("description"):
        meta["description"] = at["description"]
    if "$ref" in s and meta:
        s = {"allOf": [s], **meta}
    else:
        s.update(meta)
    if at.get("reference"):
        r = at["reference"]
        s["x-reference"] = f"{r['source']}:{r['entity']}.{r['attribute']}"
    if at.get("extension"):
        s["x-extension"] = True
    if at.get("referenceOnly"):
        s["x-reference-only"] = True
    if at.get("introducedIn"):
        s["x-introduced-in"] = at["introducedIn"]
    if at.get("sources"):
        s["x-sources"] = sorted(set(at["sources"]))
    return s


def to_openapi(model, ccfg):
    schemas = {}
    for e in model["entities"]:
        if e["kind"] == "enum":
            s = {"type": "string", "enum": e["enumValues"]}
        else:
            s = {"type": "object", "properties": {a["name"]: attr_schema(a) for a in e["attributes"]}}
            req = [a["name"] for a in e["attributes"] if a.get("required")]
            if req:
                s["required"] = req
        if e.get("description"):
            s["description"] = e["description"]
        s["x-domain"] = e.get("domain")
        s["x-approach"] = e.get("approach")
        if e.get("reference"):
            s["x-reference-entity"] = f"{e['reference']['source']}:{e['reference']['entity']}"
        s["x-introduced-in"] = e.get("introducedIn")
        s["x-status"] = e.get("status")
        if e.get("sourceSchemas"):
            s["x-source-schemas"] = e["sourceSchemas"]
        schemas[e["name"]] = s
    schemas["ProblemDetails"] = {"type": "object", "description": "RFC 7807 problem details.",
                                 "properties": {"type": {"type": "string"}, "title": {"type": "string"},
                                                "status": {"type": "integer", "format": "int32"},
                                                "detail": {"type": "string"}, "instance": {"type": "string"},
                                                "errors": {"type": "object", "additionalProperties": {"type": "array", "items": {"type": "string"}}}}}

    def ref_for(m):
        if m is None:
            return None
        if m.get("problem"):
            s = {"$ref": "#/components/schemas/ProblemDetails"}
        elif m.get("page"):
            name = f"{m['page']}Page"
            if name not in schemas:
                props = {m["itemsProperty"] or "items": {"type": "array", "items": {"$ref": f"#/components/schemas/{m['page']}"}}}
                for k, v in (m.get("other") or {}).items():
                    t, f, _, _ = parse_type(v.get("type", "integer"), v.get("format", ""))
                    props[k] = {"type": t, **({"format": f} if f in PRIM_FORMATS else {})}
                schemas[name] = {"type": "object", "description": f"A page of {m['page']}.", "properties": props,
                                 "x-approach": "envelope"}
            s = {"$ref": f"#/components/schemas/{name}"}
        elif m.get("entity"):
            s = {"$ref": f"#/components/schemas/{m['entity']}"}
        elif m.get("inline"):
            s = copy.deepcopy(m["inline"])
        elif m.get("primitive"):
            t, f, _, _ = parse_type(m["primitive"])
            s = {"type": t if t in ("string", "integer", "number", "boolean") else "string", **({"format": f} if f in PRIM_FORMATS else {})}
        else:
            return None
        return {"type": "array", "items": s} if m.get("isArray") else s

    paths, tags, sec = {}, set(), {}
    status_text = {"200": "OK", "201": "Created", "202": "Accepted", "204": "No Content", "400": "Bad Request",
                   "401": "Unauthorized", "403": "Forbidden", "404": "Not Found", "409": "Conflict", "422": "Unprocessable Entity"}
    for ep in model["endpoints"]:
        c = ep["contract"] or {}
        op = {"operationId": ep["operationId"], "tags": [ep["domain"] or "Default"], "summary": ep["summary"]}
        tags.add(ep["domain"] or "Default")
        params = []
        for p in c.get("parameters", []):
            if p.get("entity"):
                sc = {"$ref": f"#/components/schemas/{p['entity']}"}
            else:
                t, f, _, _ = parse_type(p.get("type", "string"), p.get("format", ""))
                sc = {"type": t if t in ("string", "integer", "number", "boolean") else "string", **({"format": f} if f in PRIM_FORMATS else {})}
            if p.get("isArray"):
                sc = {"type": "array", "items": sc}
            params.append({"name": p["name"], "in": p["in"], "required": True if p["in"] == "path" else p["required"], "schema": sc})
        for pp in re.findall(r"\{([^}]+)\}", ep["path"]):
            if not any(x["name"] == pp and x["in"] == "path" for x in params):
                params.append({"name": pp, "in": "path", "required": True, "schema": {"type": "string"}})
        params = [p for p in params if p["in"] != "path" or "{" + p["name"] + "}" in ep["path"]]
        if params:
            op["parameters"] = params
        if c.get("request"):
            rs = ref_for(c["request"])
            if rs:
                op["requestBody"] = {"required": True, "content": {c["request"].get("contentType", "application/json"): {"schema": rs}}}
        responses = {}
        for code, m in sorted((c.get("responses") or {}).items()):
            r = {"description": status_text.get(code, "Response")}
            rs = ref_for(m)
            if rs and code != "204":
                r["content"] = {"application/json": {"schema": rs}}
            responses[code] = r
        op["responses"] = responses or {"200": {"description": "OK"}}
        schemes = c.get("security") or []
        if "bearer" in schemes or "functionKey" in schemes:
            sec["bearer"] = {"type": "http", "scheme": "bearer", "bearerFormat": "JWT"}
            op["security"] = [{"bearer": []}]
        op["x-capability"] = ep["capability"]
        op["x-source-operations"] = [s["id"] for s in ep["sources"]]
        if ep["variants"]:
            op["x-contract-variants"] = [v["source"] for v in ep["variants"]]
        paths.setdefault(ep["path"], {})[ep["method"].lower()] = op
    spec = {
        "openapi": "3.0.3",
        "info": {"title": model["title"], "version": model["version"],
                 "description": f"Canonical API of region {model['region']}. Generated from approved alignment decisions; do not edit by hand.",
                 "x-region": model["region"], "x-status": model["status"],
                 "x-baseline": model["baseline"], "x-references": model["references"], "x-approved-by": model["reviewers"]},
        "servers": ccfg.get("servers") or [{"url": "https://{host}", "variables": {"host": {"default": f"api.{model['region'].lower()}.example"}}}],
        "tags": [{"name": t} for t in sorted(tags)],
        "paths": dict(sorted(paths.items())),
        "components": {"schemas": dict(sorted(schemas.items()))},
    }
    if sec:
        spec["components"]["securitySchemes"] = sec
    return spec


# ----------------------------------------------------------------------------- diff / outputs
def diff_models(old, new, label):
    def idx(m):
        out = {}
        for e in m.get("entities", []):
            attrs = e["attributes"] if isinstance(e["attributes"], list) else list(e["attributes"].values())
            out[e["name"]] = {a["name"]: (a.get("refEntity") or a.get("type"), bool(a.get("isCollection"))) for a in attrs}
            if e.get("kind") == "enum":
                out[e["name"]] = {v: ("value", False) for v in e.get("enumValues", [])}
        return out
    o, n = idx(old), idx(new)
    ch = {"against": label, "addedEntities": sorted(set(n) - set(o)), "removedEntities": sorted(set(o) - set(n)),
          "addedAttributes": [], "removedAttributes": [], "changedAttributes": []}
    for en in set(o) & set(n):
        for an in set(n[en]) - set(o[en]):
            ch["addedAttributes"].append(f"{en}.{an}")
        for an in set(o[en]) - set(n[en]):
            ch["removedAttributes"].append(f"{en}.{an}")
        for an in set(o[en]) & set(n[en]):
            if o[en][an] != n[en][an]:
                ch["changedAttributes"].append(f"{en}.{an}: {o[en][an]} -> {n[en][an]}")
    for k in ("addedAttributes", "removedAttributes", "changedAttributes"):
        ch[k].sort()
    ch["breaking"] = bool(ch["removedEntities"] or ch["removedAttributes"] or ch["changedAttributes"])
    return ch


def write_all(out, model, spec):
    (out / "canonical-model.json").write_text(json.dumps(model, indent=1), encoding="utf-8")
    with open(out / "canonical-model.yaml", "w", encoding="utf-8") as f:
        f.write(f"# {model['title']} - canonical model v{model['version']} ({model['status']}). Same content as canonical-model.json.\n")
        yaml.dump(model, f, Dumper=NoAlias, sort_keys=False, allow_unicode=True, width=120)
    with open(out / "canonical-openapi.yaml", "w", encoding="utf-8") as f:
        yaml.dump(spec, f, Dumper=NoAlias, sort_keys=False, allow_unicode=True, width=120)
    (out / "canonical-openapi.json").write_text(json.dumps(spec, indent=1), encoding="utf-8")
    (out / "source-to-canonical-mapping.json").write_text(json.dumps({"region": model["region"], "version": model["version"],
                                                                      "rows": model["lineage"]}, indent=1), encoding="utf-8")
    write_md(out, model)
    write_xlsx(out, model)
    from spec_viewer import describe, build as build_viewer
    build_viewer([describe(out / "canonical-openapi.yaml", spec, out), describe(out / "canonical-model.json", model, out)],
                 out / "canonical-viewer.html", f"{model['title']} v{model['version']} ({model['status']})")


def write_md(out, model):
    s = model["summary"]
    L = [f"# {model['title']} - canonical model v{model['version']} ({model['status']})", "",
         f"Region {model['region']}" + (f", baseline {model['baseline']['region']} v{model['baseline']['version']}" if model["baseline"] else "")
         + ". References: " + "; ".join(f"{r['source']} {r['name']} {r['version']}" for r in model["references"]),
         f"Reviewers: {', '.join(model['reviewers']) or '-'}", "",
         f"{s['entities']} entities ({s['baselineEntities']} from baseline, {s['newEntities']} new), {s['attributes']} attributes "
         f"({s['extensions']} extensions), {s['endpoints']} endpoints, {s['relations']} relations.", ""]
    if model["pending"]:
        L += ["## Pending review (not in the model)", ""] + [f"- {p.get('entity') or p.get('endpoint')}: {p['state']}" for p in model["pending"]] + [""]
    L += ["## Entities", ""]
    for e in model["entities"]:
        ref = f" ← {e['reference']['source']}:{e['reference']['entity']}" if e.get("reference") else ""
        L += [f"### {e['name']}{ref}", "", f"{e.get('description', '')}", "",
              f"Approach: {e['approach']} · status: {e['status']} · domain: {e.get('domain') or '-'} · sources: {', '.join(e.get('sourceSchemas', [])) or '-'}", ""]
        if e["kind"] == "enum":
            L += ["Values: " + ", ".join(e["enumValues"]), ""]
            continue
        L += ["| Attribute | Type | Required | Reference | Extension | Description |", "|---|---|---|---|---|---|"]
        for a in e["attributes"]:
            t = (("list of " if a.get("isCollection") else "") + a["refEntity"]) if a.get("refEntity") else a["type"] + (f"({a['format']})" if a.get("format") else "")
            r = f"{a['reference']['entity']}.{a['reference']['attribute']}" if a.get("reference") else ""
            L.append(f"| {a['name']} | {t} | {'yes' if a.get('required') else ''} | {r} | {'yes' if a.get('extension') else ''} | {a.get('description', '')} |")
        L.append("")
    L += ["## Endpoints", "", "| Method | Path | operationId | Sources |", "|---|---|---|---|"]
    for ep in model["endpoints"]:
        L.append(f"| {ep['method']} | `{ep['path']}` | {ep['operationId']} | {', '.join(s['id'] for s in ep['sources'])} |")
    (out / "canonical-model.md").write_text("\n".join(L) + "\n", encoding="utf-8")
    ch = model["changes"]
    C = [f"# Changelog - {model['region']} canonical v{model['version']}", "", f"Generated {model.get('generatedAt', '')}", ""]
    for blk in [model["changes"]] + ([model["changesVsBaseline"]] if model.get("changesVsBaseline") else []):
        if "against" not in blk:
            C += [f"- {blk.get('note')}", ""]
            continue
        C += [f"## Compared with {blk['against']}" + (" - BREAKING" if blk["breaking"] else ""), ""]
        for k in ("addedEntities", "removedEntities", "addedAttributes", "removedAttributes", "changedAttributes"):
            if blk[k]:
                C += [f"**{k}**: " + ", ".join(blk[k]), ""]
    del ch
    (out / "CHANGELOG.md").write_text("\n".join(C) + "\n", encoding="utf-8")


def write_xlsx(out, model):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    H = PatternFill("solid", fgColor="1F3A5F")

    def sheet(title, headers, rows):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for r in rows:
            ws.append([("" if v is None else (", ".join(map(str, v)) if isinstance(v, (list, tuple)) else v)) for v in r])
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=i)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), H
            longest = max([len(str(h))] + [len(str(ws.cell(row=r, column=i).value or "")) for r in range(2, min(ws.max_row, 200) + 1)])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, longest + 2), 60)
        ws.freeze_panes = "B2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
    ws = wb.active
    ws.title = "Summary"
    for k, v in [("Title", model["title"]), ("Region", model["region"]), ("Version", model["version"]), ("Status", model["status"]),
                 ("Baseline", f"{model['baseline']['region']} v{model['baseline']['version']}" if model["baseline"] else "-"),
                 ("References", "; ".join(f"{r['source']} {r['name']} {r['version']}" for r in model["references"])),
                 ("Reviewers", ", ".join(model["reviewers"])), ("Generated", model.get("generatedAt", ""))] + list(model["summary"].items()):
        ws.append([k, str(v)])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 22, 90
    sheet("Entities", ["Entity", "Kind", "Domain", "Approach", "Status", "Reference", "Introduced in", "Attributes", "Source schemas",
                       "Approved by", "Description"],
          [[e["name"], e["kind"], e.get("domain"), e.get("approach"), e.get("status"),
            f"{e['reference']['source']}:{e['reference']['entity']}" if e.get("reference") else "", e.get("introducedIn"),
            len(e["attributes"]) if e["kind"] != "enum" else len(e["enumValues"]), e.get("sourceSchemas", []),
            e.get("approvedBy", []), e.get("description", "")] for e in model["entities"]])
    rows = []
    for e in model["entities"]:
        if e["kind"] == "enum":
            for v in e["enumValues"]:
                rows.append([e["name"], v, "code value", "", "", "", "yes" if v in (e.get("extensionValues") or []) else "", "", ""])
            continue
        for a in e["attributes"]:
            t = (("list of " if a.get("isCollection") else "") + a["refEntity"]) if a.get("refEntity") else a["type"] + (f"({a['format']})" if a.get("format") else "")
            rows.append([e["name"], a["name"], t, "yes" if a.get("required") else "",
                         f"{a['reference']['source']}:{a['reference']['entity']}.{a['reference']['attribute']}" if a.get("reference") else "",
                         a.get("introducedIn", ""), "yes" if a.get("extension") else "", a.get("sources", []), a.get("description", "")])
    sheet("Attributes", ["Entity", "Attribute", "Type", "Required", "Reference attribute", "Introduced in", "Extension", "Sources", "Description"], rows)
    sheet("Relations", ["From", "Relation", "To", "Via", "Cardinality"],
          [[r["from"], r["type"], r["to"], r["via"], r["cardinality"]] for r in model["relations"]])
    sheet("Endpoints", ["Method", "Canonical path", "operationId", "Domain", "Capability", "Source operations", "Contract variants"],
          [[e["method"], e["path"], e["operationId"], e["domain"], e["capability"], [s["id"] for s in e["sources"]],
            [v["source"] for v in e["variants"]]] for e in model["endpoints"]])
    sheet("Source_Mapping", ["Region", "Application", "Source schemas", "Source attribute", "App entity", "Regional entity",
                             "Canonical entity", "Canonical attribute", "Reference attribute", "Transformation", "Note"],
          [[r["region"], r["appId"], r["sourceSchemas"], r["sourceAttribute"], r["appEntity"], r["regionalEntity"], r["canonicalEntity"],
            r["canonicalAttribute"], r["referenceAttribute"], r["transformation"], r["note"]] for r in model["lineage"]])
    chg = []
    for blk in [model["changes"]] + ([model["changesVsBaseline"]] if model.get("changesVsBaseline") else []):
        if "against" not in blk:
            continue
        for k in ("addedEntities", "removedEntities", "addedAttributes", "removedAttributes", "changedAttributes"):
            for x in blk[k]:
                chg.append([blk["against"], k, x])
    sheet("Changes", ["Compared with", "Change", "Item"], chg)
    sheet("Excluded", ["Kind", "Item", "State / reason"],
          [["entity", p.get("entity") or p.get("endpoint"), p["state"]] for p in model["pending"]] +
          [["rejected entity", r["entity"], f"{r.get('reviewer')}: {r.get('comment')}"] for r in model["rejected"]] +
          [["endpoint", x["id"], x["reason"]] for x in model["excludedEndpoints"]])
    wb.save(out / "canonical-model.xlsx")


if __name__ == "__main__":
    main()
