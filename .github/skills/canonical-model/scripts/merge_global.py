#!/usr/bin/env python3
"""
Merge the approved canonical models of all regions into ONE global canonical model + ONE global OpenAPI spec,
keeping for every entity, attribute, code value and endpoint which regions have it, which regions actually use it
(have source schemas / operations there), which applications it comes from, and whether it is ACORD-based or custom.

Config: api-catalog/global.yaml (optional - without it every canonical/<REGION>/ is picked up, latest release first)
  version: 1.0.0
  title: Enterprise Canonical Insurance API
  outputDir: canonical/GLOBAL
  requireReleased: true          # only released region versions may enter the global model
  regions:                       # optional explicit list / pinned versions
    - {region: EU, model: canonical/EU/releases/1.0.0/canonical-model.json}
    - {region: UK, model: canonical/UK/releases/1.0.0/canonical-model.json}

Outputs (canonical/GLOBAL/):
  global-canonical-model.json / .yaml     merged model with region / application / origin metadata
  global-canonical-openapi.yaml / .json   one OpenAPI spec (x-regions, x-applications, x-origin on every item)
  global-source-mapping.json              all regions' source-to-canonical lineage
  global-canonical.xlsx                   Entities, Attributes, Region_Matrix, Endpoints, Source_Mapping, Conflicts
  global-canonical.html                   explorer: filter by region / application / domain / origin / API, entity and
                                          endpoint views, and the filtered OpenAPI spec live as YAML or JSON
Usage:
  python merge_global.py --catalog api-catalog [--config api-catalog/global.yaml]
"""
from __future__ import annotations

import argparse
import copy
import datetime as dt
import html
import json
import re
import sys
from pathlib import Path

import yaml

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from build_canonical import to_openapi, NoAlias  # noqa: E402


def vkey(v: str):
    return tuple(int(x) if x.isdigit() else x for x in re.split(r"[.\-]", v))


def discover(catalog: Path, require_released: bool) -> list:
    out = []
    root = catalog / "canonical"
    for d in sorted(p for p in root.iterdir() if p.is_dir() and p.name != "GLOBAL") if root.exists() else []:
        rel = d / "releases"
        versions = sorted([x for x in rel.iterdir() if (x / "canonical-model.json").exists()], key=lambda x: vkey(x.name)) if rel.exists() else []
        if versions:
            out.append({"region": d.name, "model": versions[-1] / "canonical-model.json", "released": True})
        elif (d / "canonical-model.json").exists() and not require_released:
            out.append({"region": d.name, "model": d / "canonical-model.json", "released": False})
        elif (d / "canonical-model.json").exists():
            out.append({"region": d.name, "model": d / "canonical-model.json", "released": False, "skip": "not released"})
    return out


def region_of(source: str) -> str:
    return source.split("/", 1)[0] if "/" in source else ""


def app_of(source: str) -> str:
    return source.split(":", 1)[0]


def origin_of(e: dict) -> str:
    if e.get("approach") == "custom":
        return "custom"
    if e.get("reference") and e["reference"].get("source") == "acord":
        return "ACORD"
    if e.get("reference") and e["reference"].get("source") == "baseline":
        return "baseline-custom"
    return "ACORD" if e.get("reference") else "custom"


def merge(inputs: list, gcfg: dict) -> dict:
    ents, conflicts, eps, lineage = {}, [], {}, []
    regions_meta = []
    for inp in inputs:
        m = inp["doc"]
        reg = m["region"]
        regions_meta.append({"region": reg, "version": m["version"], "status": m["status"], "released": inp["released"],
                             "baseline": m.get("baseline"), "references": m.get("references", []), "reviewers": m.get("reviewers", []),
                             "file": str(inp["model"])})
        for e in m["entities"]:
            g = ents.get(e["name"])
            srcs = e.get("sourceSchemas", [])
            used = sorted({region_of(s) for s in srcs if region_of(s)})
            apps = sorted({app_of(s) for s in srcs})
            if not g:
                g = ents[e["name"]] = {
                    "name": e["name"], "kind": e["kind"], "domain": e.get("domain"), "domains": list(e.get("domains") or []),
                    "description": e.get("description", ""), "origin": origin_of(e), "reference": e.get("reference"),
                    "introducedIn": e.get("introducedIn") or reg, "availableIn": [], "usedBy": [], "applications": [],
                    "approachByRegion": {}, "attributes": {}, "enumValues": [], "valueRegions": {}, "sourceSchemas": [],
                    "approvedBy": []}
            if reg not in g["availableIn"]:
                g["availableIn"].append(reg)
            g["usedBy"] = sorted(set(g["usedBy"]) | set(used))
            g["applications"] = sorted(set(g["applications"]) | set(apps))
            g["sourceSchemas"] = sorted(set(g["sourceSchemas"]) | set(srcs))
            g["approachByRegion"][reg] = e.get("approach")
            g["approvedBy"] = sorted(set(g["approvedBy"]) | set(e.get("approvedBy") or []))
            for d in e.get("domains") or []:
                if d and d not in g["domains"]:
                    g["domains"].append(d)
            if not g["description"] and e.get("description"):
                g["description"] = e["description"]
            if g["kind"] != e["kind"]:
                conflicts.append({"kind": "entity-kind", "item": e["name"], "detail": f"{g['kind']} vs {e['kind']} in {reg}"})
            for v in e.get("enumValues") or []:
                if v not in g["enumValues"]:
                    g["enumValues"].append(v)
                g["valueRegions"].setdefault(v, [])
                if reg not in g["valueRegions"][v]:
                    g["valueRegions"][v].append(reg)
            for a in e.get("attributes", []):
                ga = g["attributes"].get(a["name"])
                a_used = sorted({region_of(s) for s in a.get("sources", []) if region_of(s)})
                if not ga:
                    ga = g["attributes"][a["name"]] = dict(copy.deepcopy(a), availableIn=[], usedBy=[], applications=[])
                    ga["sources"] = []
                else:
                    sig_old = (ga.get("refEntity") or ga.get("type"), bool(ga.get("isCollection")), ga.get("format") or "")
                    sig_new = (a.get("refEntity") or a.get("type"), bool(a.get("isCollection")), a.get("format") or "")
                    if sig_old != sig_new:
                        conflicts.append({"kind": "attribute-type", "item": f"{e['name']}.{a['name']}",
                                          "detail": f"{sig_old} ({', '.join(ga['availableIn'])}) vs {sig_new} ({reg})"})
                    if not ga.get("description") and a.get("description"):
                        ga["description"] = a["description"]
                if reg not in ga["availableIn"]:
                    ga["availableIn"].append(reg)
                ga["usedBy"] = sorted(set(ga["usedBy"]) | set(a_used))
                ga["sources"] = sorted(set(ga["sources"]) | set(a.get("sources", [])))
                ga["applications"] = sorted({app_of(s) for s in ga["sources"]})
        for ep in m.get("endpoints", []):
            k = (ep["method"], ep["path"])
            g = eps.get(k)
            src_ids = [s["id"] for s in ep["sources"]]
            if not g:
                g = eps[k] = dict(copy.deepcopy(ep), regions=[], applications=[], sources=[])
            if reg not in g["regions"]:
                g["regions"].append(reg)
            known = {s["id"] for s in g["sources"]}
            g["sources"] += [s for s in ep["sources"] if s["id"] not in known]
            g["applications"] = sorted({s["id"].split(":")[0] for s in g["sources"]})
            if json.dumps(ep.get("contract"), sort_keys=True) != json.dumps(g.get("contract"), sort_keys=True) and \
                    not any(v.get("region") == reg for v in g.get("variants", [])):
                g.setdefault("variants", []).append({"region": reg, "source": ",".join(src_ids), "contract": ep.get("contract")})
        for r in m.get("lineage", []):
            lineage.append(r)
    # operationId uniqueness across regions
    seen = {}
    for k, ep in sorted(eps.items(), key=lambda kv: (kv[0][1], kv[0][0])):
        oid = ep["operationId"]
        if oid in seen and seen[oid] != k:
            new = oid + "".join(ep["regions"])
            conflicts.append({"kind": "operationId", "item": oid, "detail": f"{seen[oid]} and {k} - renamed to {new}"})
            ep["operationId"] = new
        seen[ep["operationId"]] = k
    entities = []
    for g in sorted(ents.values(), key=lambda x: x["name"]):
        g2 = dict(g)
        g2["attributes"] = list(g["attributes"].values())
        g2["status"] = "approved"
        g2["approach"] = "global"
        entities.append(g2)
    relations = []
    for e in entities:
        for a in e["attributes"]:
            if a.get("refEntity") and a["refEntity"] in ents:
                k = ents[a["refEntity"]]["kind"]
                relations.append({"from": e["name"], "to": a["refEntity"], "via": a["name"],
                                  "type": "uses-code-list" if k == "enum" else ("has-many" if a.get("isCollection") else "has-one"),
                                  "cardinality": "0..*" if a.get("isCollection") else ("1" if a.get("required") else "0..1"),
                                  "regions": a["availableIn"]})
    regions = [r["region"] for r in regions_meta]
    return {
        "modelType": "canonical-global", "region": "GLOBAL", "version": str(gcfg.get("version", "1.0.0")),
        "title": gcfg.get("title") or "Enterprise Canonical Insurance API", "status": "approved",
        "regions": regions_meta, "references": [json.loads(x) for x in sorted({json.dumps(r, sort_keys=True) for m in regions_meta for r in m["references"]})],
        "reviewers": sorted({x for m in regions_meta for x in m["reviewers"]}),
        "entities": entities, "relations": relations, "endpoints": [eps[k] for k in sorted(eps, key=lambda k: (k[1], k[0]))],
        "lineage": lineage, "conflicts": conflicts,
        "summary": {"regions": len(regions), "entities": len(entities), "acordEntities": sum(1 for e in entities if e["origin"] == "ACORD"),
                    "customEntities": sum(1 for e in entities if e["origin"] != "ACORD"),
                    "attributes": sum(len(e["attributes"]) for e in entities), "endpoints": len(eps),
                    "sharedByAllRegions": sum(1 for e in entities if len(e["availableIn"]) == len(regions)),
                    "conflicts": len(conflicts), "lineageRows": len(lineage)},
    }


def global_openapi(model: dict, gcfg: dict) -> dict:
    tmp = {"title": model["title"], "version": model["version"], "region": "GLOBAL", "status": model["status"],
           "baseline": None, "references": model["references"], "reviewers": model["reviewers"],
           "entities": model["entities"], "endpoints": model["endpoints"]}
    spec = to_openapi(tmp, gcfg)
    spec["info"]["description"] = ("Global canonical API merged from the approved canonical models of regions "
                                   + ", ".join(f"{r['region']} v{r['version']}" for r in model["regions"]) + ". Do not edit by hand.")
    spec["info"]["x-regions"] = [{k: r[k] for k in ("region", "version", "released")} for r in model["regions"]]
    by = {e["name"]: e for e in model["entities"]}
    for n, s in spec["components"]["schemas"].items():
        e = by.get(n)
        if not e:
            continue
        s["x-origin"] = e["origin"]
        s["x-available-in"] = e["availableIn"]
        s["x-used-by-regions"] = e["usedBy"]
        s["x-applications"] = e["applications"]
        for a in e["attributes"]:
            ps = (s.get("properties") or {}).get(a["name"])
            if ps is not None:
                ps["x-available-in"] = a["availableIn"]
                if a["usedBy"]:
                    ps["x-used-by-regions"] = a["usedBy"]
        if e["kind"] == "enum" and len(e["availableIn"]) > 1:
            s["x-value-regions"] = e["valueRegions"]
    eps = {(e["method"].lower(), e["path"]): e for e in model["endpoints"]}
    for p, item in spec["paths"].items():
        for m, op in item.items():
            e = eps.get((m, p))
            if e:
                op["x-regions"] = e["regions"]
                op["x-applications"] = e["applications"]
    return spec


def write_xlsx(out: Path, model: dict):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    H = PatternFill("solid", fgColor="1F3A5F")
    regs = [r["region"] for r in model["regions"]]

    def sheet(title, headers, rows):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for r in rows:
            ws.append([("" if v is None else (", ".join(map(str, v)) if isinstance(v, (list, tuple)) else v)) for v in r])
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=i)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), H
            longest = max([len(str(h))] + [len(str(ws.cell(row=r, column=i).value or "")) for r in range(2, min(ws.max_row, 200) + 1)])
            ws.column_dimensions[get_column_letter(i)].width = min(max(9, longest + 2), 60)
        ws.freeze_panes = "B2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
    ws = wb.active
    ws.title = "Summary"
    for k, v in [("Title", model["title"]), ("Version", model["version"]), ("Generated", model.get("generatedAt", ""))] + \
            [(f"Region {r['region']}", f"v{r['version']} {'released' if r['released'] else 'NOT released'}") for r in model["regions"]] + \
            list(model["summary"].items()):
        ws.append([k, str(v)])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 24, 80
    sheet("Region_Matrix", ["Entity", "Kind", "Origin", "Domain"] + [f"{r} available" for r in regs] + [f"{r} used" for r in regs] + ["Applications"],
          [[e["name"], e["kind"], e["origin"], e.get("domain")] + [("yes" if r in e["availableIn"] else "") for r in regs] +
           [("yes" if r in e["usedBy"] else "") for r in regs] + [e["applications"]] for e in model["entities"]])
    sheet("Entities", ["Entity", "Kind", "Origin", "Reference", "Introduced in", "Available in", "Used by regions", "Applications",
                       "Attributes", "Description"],
          [[e["name"], e["kind"], e["origin"], f"{e['reference']['source']}:{e['reference']['entity']}" if e.get("reference") else "",
            e["introducedIn"], e["availableIn"], e["usedBy"], e["applications"],
            len(e["attributes"]) if e["kind"] != "enum" else len(e["enumValues"]), e.get("description", "")] for e in model["entities"]])
    rows = []
    for e in model["entities"]:
        if e["kind"] == "enum":
            for v in e["enumValues"]:
                rows.append([e["name"], v, "code value", "", "", e["valueRegions"].get(v, []), "", "", ""])
            continue
        for a in e["attributes"]:
            t = (("list of " if a.get("isCollection") else "") + a["refEntity"]) if a.get("refEntity") else a["type"] + (f"({a['format']})" if a.get("format") else "")
            rows.append([e["name"], a["name"], t, f"{a['reference']['entity']}.{a['reference']['attribute']}" if a.get("reference") else "",
                         "yes" if a.get("extension") else "", a["availableIn"], a["usedBy"], a["applications"], a.get("description", "")])
    sheet("Attributes", ["Entity", "Attribute", "Type", "Reference attribute", "Extension", "Available in", "Used by regions", "Applications", "Description"], rows)
    sheet("Endpoints", ["Method", "Path", "operationId", "Domain", "Regions", "Applications", "Source operations"],
          [[e["method"], e["path"], e["operationId"], e.get("domain"), e["regions"], e["applications"], [s["id"] for s in e["sources"]]]
           for e in model["endpoints"]])
    sheet("Source_Mapping", ["Region", "Application", "Source schemas", "Source attribute", "Canonical entity", "Canonical attribute",
                             "Reference attribute", "Transformation", "Note"],
          [[r["region"], r["appId"], r["sourceSchemas"], r["sourceAttribute"], r["canonicalEntity"], r["canonicalAttribute"],
            r["referenceAttribute"], r["transformation"], r["note"]] for r in model["lineage"]])
    sheet("Conflicts", ["Kind", "Item", "Detail"], [[c["kind"], c["item"], c["detail"]] for c in model["conflicts"]])
    wb.save(out / "global-canonical.xlsx")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--catalog", default="api-catalog")
    ap.add_argument("--config")
    a = ap.parse_args()
    catalog = Path(a.catalog).resolve()
    cfg_p = Path(a.config) if a.config else catalog / "global.yaml"
    gcfg = (yaml.safe_load(cfg_p.read_text(encoding="utf-8")) or {}) if cfg_p.exists() else {}
    require_released = bool(gcfg.get("requireReleased", True))
    if gcfg.get("regions"):
        inputs = [{"region": r["region"], "model": (catalog / r["model"]).resolve(), "released": "releases" in Path(r["model"]).parts}
                  for r in gcfg["regions"]]
    else:
        inputs = discover(catalog, require_released)
    skipped = [i for i in inputs if i.get("skip")]
    inputs = [i for i in inputs if not i.get("skip")]
    for i in inputs:
        i["doc"] = json.loads(Path(i["model"]).read_text(encoding="utf-8"))
        if i["doc"].get("status") != "approved":
            sys.exit(f"{i['model']}: status {i['doc'].get('status')} - only approved region models enter the global model")
        if require_released and not i["released"]:
            sys.exit(f"{i['model']}: not a release - release the region first or set requireReleased: false")
    if not inputs:
        sys.exit("No approved region canonical models found under canonical/<REGION>/")
    out = (catalog / gcfg.get("outputDir", "canonical/GLOBAL")).resolve()
    out.mkdir(parents=True, exist_ok=True)
    model = merge(inputs, gcfg)
    model["generatedAt"] = dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")
    model["skippedRegions"] = [{"region": s["region"], "reason": s["skip"]} for s in skipped]
    for r in model["regions"]:
        r["file"] = str(Path(r["file"]).relative_to(catalog)) if str(r["file"]).startswith(str(catalog)) else r["file"]
    spec = global_openapi(model, gcfg)
    from sensitive_scan import find_policy_file, load_policy, redact_obj
    pol = load_policy(find_policy_file(out))
    model, _ = redact_obj(model, pol)
    spec, _ = redact_obj(spec, pol)
    (out / "global-canonical-model.json").write_text(json.dumps(model, indent=1), encoding="utf-8")
    with open(out / "global-canonical-model.yaml", "w", encoding="utf-8") as f:
        f.write(f"# {model['title']} v{model['version']} - same content as global-canonical-model.json\n")
        yaml.dump(model, f, Dumper=NoAlias, sort_keys=False, allow_unicode=True, width=120)
    with open(out / "global-canonical-openapi.yaml", "w", encoding="utf-8") as f:
        yaml.dump(spec, f, Dumper=NoAlias, sort_keys=False, allow_unicode=True, width=120)
    (out / "global-canonical-openapi.json").write_text(json.dumps(spec, indent=1), encoding="utf-8")
    (out / "global-source-mapping.json").write_text(json.dumps({"rows": model["lineage"]}, indent=1), encoding="utf-8")
    write_xlsx(out, model)
    tpl = (HERE.parent / "assets" / "global-canonical.template.html").read_text(encoding="utf-8")
    payload = json.dumps({"model": {k: v for k, v in model.items() if k != "lineage"}, "spec": spec}, separators=(",", ":")).replace("</", "<\\/")
    (out / "global-canonical.html").write_text(tpl.replace("{{TITLE}}", html.escape(f"{model['title']} v{model['version']}"))
                                               .replace("/*__DATA__*/null", payload), encoding="utf-8")
    print(json.dumps({"output": str(out), "regions": [f"{r['region']} v{r['version']}{'' if r['released'] else ' (not released)'}" for r in model["regions"]],
                      "skipped": model["skippedRegions"], **model["summary"]}, indent=2))


if __name__ == "__main__":
    main()
