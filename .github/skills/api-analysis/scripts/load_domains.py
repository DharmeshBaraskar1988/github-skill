#!/usr/bin/env python3
"""
Convert the business domain catalogue (Excel) into domains.json.

Two layouts are accepted (first sheet, or --sheet). Header names are matched
case-insensitively and loosely; only "Domain" is mandatory.

Grouped layout (business capability map) - domain name only on its first row, merged cells allowed:

  | Domain | Domain Description | SubCapability | SubCapability Description | Business Outcomes | Domain Keywords | SubCapability Keywords |
  | Claims Management | ... | FNOL | Captures First Notification of Loss ... | Fair Claims Handling | claim, claimant, loss | fnol, register claim |
  |                   |     | Settlement | ...                                | Loss Control         |                       | approve, settle      |
  | (Including Bind)  |     | ...  <- a "(...)" cell under a domain name is a scope note: its terms become domain keywords

  Empty Domain cells belong to the domain above. Business Outcomes are collected per domain.
  Without keyword columns, keywords are derived from the names and scope notes (marked keywordsDerived).

Flat layout:

  | Domain | Capability | Description | Keywords | Region | Application |
  |--------|------------|-------------|----------|--------|-------------|
  | Claims |            | Claims mgmt | claim, fnol, loss, settlement | | |
  | Claims | Register Claim | ...     | create claim, fnol           | | |

* A row with a Domain and no Capability describes the domain itself.
* A row with a Capability adds that capability to the domain.
* Keywords: comma / semicolon / newline separated. They drive the mapping
  of endpoints to domains, so list the nouns your APIs actually use.
* Region / Application (optional): restrict a row to some regions/apps
  (comma separated, empty = everywhere).

Usage:
  python load_domains.py --xlsx domains.xlsx --out api-catalog/domains.json
  python load_domains.py --template domains-template.xlsx      # write an example workbook
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl")

HEADER_ALIASES = {
    "domain": ["domain", "business domain", "domain name", "bounded context"],
    "subdomain": ["sub domain", "subdomain", "sub-domain"],
    "capability": ["capability", "business capability", "capability name", "function",
                   "subcapability", "sub capability", "sub-capability", "sub capabilities", "subcapabilities"],
    "description": ["description", "definition", "details"],
    "domain_description": ["domain description", "domain definition"],
    "capability_description": ["subcapability description", "sub capability description", "sub-capability description",
                               "capability description"],
    "outcomes": ["business outcomes", "business outcome", "outcomes", "benefits"],
    "keywords": ["keywords", "keyword", "synonyms", "terms", "aliases"],
    "domain_keywords": ["domain keywords", "domain keyword", "domain synonyms"],
    "capability_keywords": ["subcapability keywords", "sub capability keywords", "sub-capability keywords",
                            "capability keywords"],
    "region": ["region", "regions"],
    "application": ["application", "applications", "app", "project", "projects", "system"],
}


def norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


GENERIC = {"management", "manage", "lifecycle", "data", "others", "other", "check", "including", "service", "services",
            "processing", "and", "the", "of", "&", "-", "s"}


def derive_keywords(*texts) -> list:
    """Keywords from names / scope notes when the Excel has none: lower-case words minus generic ones."""
    out = []
    for t in texts:
        for w in re.split(r"[^A-Za-z0-9-]+", str(t or "")):
            w = w.strip("-").lower()
            if len(w) > 1 and w not in GENERIC and w not in out:
                out.append(w)
    return out


def split_list(s) -> list:
    return [x.strip() for x in re.split(r"[,;\n|]", str(s or "")) if x.strip()]


def map_headers(row) -> dict:
    idx = {}
    for i, cell in enumerate(row):
        h = norm(cell).lower()
        for key, aliases in HEADER_ALIASES.items():
            if h in aliases and key not in idx:
                idx[key] = i
    return idx


def load(xlsx: Path, sheet: str | None) -> dict:
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header_row, idx = None, {}
    for r_i, row in enumerate(rows[:20]):
        idx = map_headers(row)
        if "domain" in idx:
            header_row = r_i
            break
    if header_row is None:
        sys.exit(f"No 'Domain' header found in the first 20 rows of sheet '{ws.title}'")

    domains: dict = {}
    issues = []
    current = None
    for n, row in enumerate(rows[header_row + 1:], start=header_row + 2):
        get = lambda k: row[idx[k]] if k in idx and idx[k] < len(row) else None  # noqa: E731
        if not any(norm(c) for c in row):
            continue
        dcell = norm(get("domain"))
        scope = ""
        if dcell.startswith("(") and dcell.endswith(")") and current:
            scope = dcell[1:-1]                      # "(KYC, Sanctions, ...)" under a domain name
        elif dcell:
            current = dcell
        if not current:
            issues.append(f"row {n}: no Domain value - skipped")
            continue
        dname = current
        d = domains.setdefault(dname, {"name": dname, "description": "", "keywords": [], "subdomains": [], "scope": [],
                                       "outcomes": [], "regions": [], "applications": [], "capabilities": [],
                                       "keywordsDerived": False})
        cap = norm(get("capability"))
        generic_desc = norm(get("description"))
        d_desc = norm(get("domain_description")) or ("" if cap else generic_desc)
        c_desc = norm(get("capability_description")) or (generic_desc if cap else "")
        generic_kws = split_list(get("keywords"))
        d_kws = split_list(get("domain_keywords")) + ([] if cap else generic_kws)
        c_kws = split_list(get("capability_keywords")) + (generic_kws if cap else [])
        regions = split_list(get("region"))
        apps = split_list(get("application"))
        sub = norm(get("subdomain"))
        if scope:
            d["scope"] += [x for x in split_list(scope) if x not in d["scope"]]
        if sub and sub not in d["subdomains"]:
            d["subdomains"].append(sub)
        if d_desc and not d["description"]:
            d["description"] = d_desc
        d["keywords"] += [k for k in d_kws if k not in d["keywords"]]
        for o in split_list(get("outcomes")):
            if o not in d["outcomes"]:
                d["outcomes"].append(o)
        if not cap:
            d["regions"] += [r for r in regions if r not in d["regions"]]
            d["applications"] += [a for a in apps if a not in d["applications"]]
        else:
            if any(c["name"].lower() == cap.lower() for c in d["capabilities"]):
                issues.append(f"row {n}: duplicate capability '{cap}' in domain '{dname}'")
                continue
            d["capabilities"].append({"name": cap, "description": c_desc, "keywords": c_kws,
                                      "keywordsDerived": not c_kws,
                                      "subdomain": sub, "regions": regions, "applications": apps, "row": n})
    for d in domains.values():
        if not d["keywords"]:
            d["keywords"] = derive_keywords(d["name"], *d["scope"])
            d["keywordsDerived"] = True
            issues.append(f"domain '{d['name']}': no keywords in the Excel - derived {d['keywords']}; add a Domain Keywords "
                          f"column with the nouns your APIs use for reliable mapping")
        else:
            d["keywords"] += [k for k in derive_keywords(*d["scope"]) if k not in d["keywords"]]
        for c in d["capabilities"]:
            if not c["keywords"]:
                c["keywords"] = derive_keywords(c["name"])
        if not d["capabilities"]:
            issues.append(f"domain '{d['name']}': no capabilities - endpoints mapped to it get proposed capabilities")
    return {"source": str(xlsx), "sheet": ws.title, "domains": list(domains.values()), "issues": issues}


def write_template(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Domains"
    ws.append(["Domain", "Capability", "Description", "Keywords", "Region", "Application"])
    data = [
        ("Claims", "", "Claims handling from first notice of loss to settlement", "claim, claims, fnol, loss, claimant, settlement, reserve", "", ""),
        ("Claims", "Register Claim", "Create a claim / first notice of loss", "create claim, register, fnol, submit", "", ""),
        ("Claims", "View Claim", "Retrieve claim details", "get claim, view, details", "", ""),
        ("Claims", "Search Claims", "List and search claims", "list claims, search, find", "", ""),
        ("Claims", "Update Claim", "Change claim data", "update claim, modify", "", ""),
        ("Claims", "Approve Claim", "Approve and settle a claim", "approve, settlement, settle", "", ""),
        ("Claims", "Close Claim", "Close or withdraw a claim", "delete claim, close, withdraw", "", ""),
        ("Claims", "Manage Claimants", "Parties claiming under a policy", "claimant, claimants", "", ""),
        ("Document Management", "", "Documents attached to business objects", "document, documents, upload, attachment, file", "", ""),
        ("Document Management", "Upload Document", "", "upload document, attach", "", ""),
        ("Document Management", "List Documents", "", "list documents", "", ""),
        ("Document Management", "Delete Document", "", "delete document, remove", "", ""),
        ("Policy", "", "Policy administration", "policy, policies, coverage, premium, endorsement, renewal, policyholder", "", ""),
        ("Policy", "Issue Policy", "", "create policy, issue, bind", "", ""),
        ("Policy", "View Policy", "", "get policy, view policy", "", ""),
        ("Policy", "Search Policies", "", "list policies, search", "", ""),
        ("Policy", "Renew Policy", "", "renew, renewal", "", ""),
        ("Policy", "Manage Policyholder", "", "policyholder, holder, insured", "", ""),
        ("Quote", "", "Quotation and pricing", "quote, quotes, quotation, price, pricing, rating", "", ""),
        ("Quote", "Create Quote", "", "create quote, quotation", "", ""),
        ("Quote", "Get Quote", "", "get quote", "", ""),
    ]
    for r in data:
        ws.append(list(r))
    for col, w in zip("ABCDEF", (24, 26, 50, 60, 12, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    readme = wb.create_sheet("README")
    for line in (__doc__ or "").splitlines():
        readme.append([line])
    readme.column_dimensions["A"].width = 110
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx")
    ap.add_argument("--sheet")
    ap.add_argument("--out")
    ap.add_argument("--template", help="write an example domain workbook to this path and exit")
    a = ap.parse_args()
    if a.template:
        write_template(Path(a.template))
        print(f"Template written: {a.template}")
        return
    if not a.xlsx or not a.out:
        ap.error("--xlsx and --out are required")
    res = load(Path(a.xlsx), a.sheet)
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    Path(a.out).write_text(json.dumps(res, indent=2), encoding="utf-8")
    print(json.dumps({"output": a.out, "domains": len(res["domains"]),
                      "capabilities": sum(len(d["capabilities"]) for d in res["domains"]),
                      "issues": res["issues"]}, indent=2))


if __name__ == "__main__":
    main()
